"""Achados no modelo de planilha "Casos Estaduais" (D112): mapeamento, resumo por UF, formato do arquivo, leitura do banco,
comando e download pelo painel."""
import io
import re
import sqlite3
import sys

import openpyxl
import pytest
from fastapi.testclient import TestClient

import main
from adaptadores.bruto_local import BrutoLocal
from adaptadores.leitura_sqlite import LeituraPainelSQLite
from adaptadores.repositorio_sqlite import RepositorioSQLite
from aplicacao import exportacao
from config import consultas as C, vocabulario as V
from dominio import casos_estaduais as CE
from painel import rotulos
from painel.app import Contexto, criar_app

UFS = {"PI": "Piauí", "SP": "São Paulo", "AC": "Acre"}
TEC = {"reconhecimento_facial": "Reconhecimento facial", "drones": "Drones"}
ST = {"em_uso": "Em uso", "piloto": "Piloto", "descartado": "Descartado (achado incorreto)"}


def _l(**kw):
    base = {"achado_id": 1, "uf": "SP", "cidade": None, "tecnologia": "reconhecimento_facial", "status_reportado": "piloto",
            "status_final": None, "verificado": 0, "fornecedor": None, "nome_solucao": None, "observacao": None,
            "url": "https://ssp.sp.gov.br/a", "dominio": "ssp.sp.gov.br", "fonte_oficial": 1, "titulo": "Título",
            "data_pub": "2025-12-05"}
    base.update(kw)
    return base


# ---------- domínio ----------
@pytest.mark.parametrize("texto, ano", [
    ("2025-12-05", 2025), ("dez. de 2024", 2024), ("2025/2026", 2025), ("05/12/1999", 1999),
    ("12/05/1899", None), ("20250", None), ("", None), (None, None), (2025, None), ("abc", None),
])
def test_ano_da_data(texto, ano):
    assert CE.ano_da_data(texto) == ano


def test_linha_de_achado_nao_verificado_usa_o_status_do_classificador():
    assert CE.linha_do_caso(_l(), UFS, TEC, ST) == {
        "ID": 1, "Categoria": "Reconhecimento facial", "Autor/Instituicao": "ssp.sp.gov.br", "Ano": 2025, "Titulo": "Título",
        "Tipo de fonte": "Notícia oficial / literatura cinzenta", "Pais/escopo": "Brasil/São Paulo",
        "Topicos e palavras-chave": "situação: Piloto", "URL": "https://ssp.sp.gov.br/a", "Prioridade de leitura": "",
        "Status de leitura": "Nao iniciado", "Notas": ""}


def test_linha_de_achado_verificado_usa_a_decisao_humana_e_leva_tudo():
    l = _l(verificado=1, status_final="em_uso", fornecedor="Empresa X", nome_solucao="Olho Vivo", observacao="confirmado",
           cidade="Campinas", fonte_oficial=0, dominio="g1.globo.com", tecnologia="drones", data_pub=None, titulo=None, url=None)
    assert CE.linha_do_caso(l, UFS, TEC, ST) == {
        "ID": 1, "Categoria": "Drones", "Autor/Instituicao": "g1.globo.com", "Ano": None, "Titulo": "",
        "Tipo de fonte": "Reportagem / literatura cinzenta", "Pais/escopo": "Brasil/São Paulo/Campinas",
        "Topicos e palavras-chave": "fornecedor: Empresa X; solução: Olho Vivo; situação: Em uso", "URL": "",
        "Prioridade de leitura": "", "Status de leitura": "Lido", "Notas": "confirmado"}


def test_rotulo_desconhecido_fica_com_a_chave():
    l = _l(tecnologia="outra_coisa", status_reportado="novo_status", uf="XX")
    linha = CE.linha_do_caso(l, UFS, TEC, ST)
    assert (linha["Categoria"], linha["Topicos e palavras-chave"], linha["Pais/escopo"]) == ("outra_coisa", "situação: novo_status", "Brasil/XX")


def test_casos_tiram_os_descartados_e_seguem_a_ordem_das_ufs():
    linhas = [_l(achado_id=5, uf="SP"), _l(achado_id=3, uf="PI", tecnologia="drones"), _l(achado_id=4, uf="PI"),
              _l(achado_id=2, uf="PI", verificado=1, status_final="descartado"), _l(achado_id=1, uf="SP", tecnologia="drones"),
              _l(achado_id=6, uf="PI", verificado=0, status_final="descartado")]          # descarte desfeito: volta a valer
    # PI antes de SP (ordem de UFS); rótulo da tecnologia (Drones antes de Reconhecimento facial); número
    assert [c["ID"] for c in CE.casos(linhas, UFS, TEC, ST)] == [3, 4, 6, 1, 5]


def test_ordem_pelo_rotulo_sem_diferenca_de_acento():
    tec = {"z_chave": "Análise preditiva", "a_chave": "Busca", "m_chave": "analise de vínculos"}
    linhas = [_l(achado_id=1, tecnologia="a_chave"), _l(achado_id=2, tecnologia="z_chave"), _l(achado_id=3, tecnologia="m_chave")]
    assert [c["Categoria"] for c in CE.casos(linhas, UFS, tec, ST)] == ["analise de vínculos", "Análise preditiva", "Busca"]
    assert CE.resumo(linhas, UFS, tec)[0]["Dados e evidencias de discriminacao"].startswith(
        "analise de vínculos (1), Análise preditiva (1), Busca (1);")


def test_verificacao_desfeita_nao_leva_a_observacao_antiga():
    linha = CE.linha_do_caso(_l(verificado=0, observacao="nota de uma verificação revogada"), UFS, TEC, ST)
    assert (linha["Status de leitura"], linha["Notas"]) == ("Nao iniciado", "")


def test_resumo_por_uf():
    linhas = [_l(achado_id=1, uf="SP"), _l(achado_id=2, uf="SP", verificado=1, status_final="em_uso"),
              _l(achado_id=3, uf="SP", tecnologia="drones"), _l(achado_id=4, uf="SP", tecnologia="drones"),
              _l(achado_id=5, uf="SP", tecnologia="drones", verificado=1, status_final="descartado"),
              _l(achado_id=6, uf="PI"), _l(achado_id=7, uf="ZZ")]
    assert CE.resumo(linhas, UFS, TEC) == [
        {"Estado/Caso": "Piauí (PI)", "Dados e evidencias de discriminacao": "Reconhecimento facial (1); 0 verificado(s) de 1"},
        {"Estado/Caso": "São Paulo (SP)", "Dados e evidencias de discriminacao": "Drones (2), Reconhecimento facial (2); 1 verificado(s) de 4"},
        {"Estado/Caso": "ZZ (ZZ)", "Dados e evidencias de discriminacao": "Reconhecimento facial (1); 0 verificado(s) de 1"}]


# ---------- arquivo ----------
def _abrir(conteudo):
    return openpyxl.load_workbook(io.BytesIO(conteudo))


def test_arquivo_tem_as_duas_abas_no_formato_do_modelo():
    linhas = [_l(achado_id=1), _l(achado_id=2, uf="PI", titulo="=HYPERLINK(\"x\")")]
    wb = _abrir(exportacao.planilha_casos_estaduais(linhas, UFS, TEC, ST))
    assert wb.sheetnames == ["Casos Estaduais", "Resumo Casos"]
    casos, resumo = wb["Casos Estaduais"], wb["Resumo Casos"]
    assert [c.value for c in casos[1]] == list(CE.COLUNAS)
    assert [c.value for c in resumo[1]] == ["Estado/Caso", "Dados e evidencias de discriminacao"]
    assert [casos.column_dimensions[l].width for l in "ABCDEFGHIJKL"] == [12, 38, 42, 12, 70, 26, 23, 70, 41, 23, 19, 12]
    assert [resumo.column_dimensions[l].width for l in "AB"] == [23, 70]
    assert (casos.freeze_panes, casos.auto_filter.ref) == ("A2", "A1:L3")
    assert (resumo.freeze_panes, resumo.auto_filter.ref) == ("A2", "A1:B3")
    assert [c.value for c in casos[2]][:2] == [2, "Reconhecimento facial"]   # PI antes de SP
    assert not str(casos["E2"].value).startswith("=")                      # regra C3: fórmula não entra como fórmula
    for c in casos[1] + resumo[1]:                                        # cabeçalho como no modelo
        assert (c.fill.fgColor.rgb, c.font.color.rgb, c.font.b, c.alignment.vertical) == ("004472C4", "00FFFFFF", True, "top")
    assert casos["A2"].fill.fgColor.rgb != "004472C4"


def test_formulas_em_qualquer_texto_nao_viram_formula():
    l = _l(verificado=1, status_final="em_uso", observacao="=1+1", fornecedor="=CMD()", nome_solucao="@SOMA(A1)",
           titulo="=HYPERLINK(1)", dominio="=x", url="=y")
    wb = _abrir(exportacao.planilha_casos_estaduais([l], UFS, TEC, ST))
    for c in wb["Casos Estaduais"][2]:
        assert not (isinstance(c.value, str) and c.value.startswith("=")), c.coordinate


def test_arquivo_sem_achados_tem_so_os_cabecalhos():
    wb = _abrir(exportacao.planilha_casos_estaduais([], UFS, TEC, ST))
    assert wb["Casos Estaduais"].max_row == 1 and wb["Casos Estaduais"].auto_filter.ref == "A1:L1"
    assert wb["Resumo Casos"].max_row == 1 and wb["Resumo Casos"].auto_filter.ref == "A1:B1"


@pytest.mark.parametrize("n, letras", [(1, "A"), (12, "L"), (26, "Z"), (27, "AA"), (52, "AZ"), (53, "BA"), (702, "ZZ"), (703, "AAA")])
def test_letra_da_coluna(n, letras):
    assert exportacao._letra_da_coluna(n) == letras


def test_planilha_sem_larguras_continua_como_antes():
    import pandas as pd
    wb = _abrir(exportacao.planilha_xlsx({"A": pd.DataFrame({"x": [1]})}))
    assert wb["A"].freeze_panes is None and wb["A"].auto_filter.ref is None


# ---------- leitura do banco (demonstração) ----------
@pytest.fixture(scope="module")
def demo(copiar_demo):
    banco, raw = copiar_demo("casos_estaduais")
    return banco, raw


def test_leitura_traz_todos_os_achados_do_relatorio(demo):
    banco, _ = demo
    linhas = LeituraPainelSQLite(banco).linhas_casos_estaduais()
    assert len(linhas) == sqlite3.connect(banco).execute("SELECT COUNT(*) FROM v_achados_revisao").fetchone()[0] > 0


def test_titulo_e_data_vem_do_resultado_mais_recente_da_mesma_pagina(demo, tmp_path):
    import shutil
    banco = tmp_path / "t.db"
    shutil.copy(demo[0], banco)
    con = sqlite3.connect(banco)
    a = con.execute("SELECT a.achado_id, k.url_id, k.uf, k.run_id FROM achados a JOIN r_classificacoes k "
                    "ON k.classificacao_id = a.classificacao_id ORDER BY a.achado_id LIMIT 1").fetchone()
    def chamada(quando):
        return con.execute("INSERT INTO chamadas (run_id, etapa, provedor, parametros_json, chave_idempotencia, requisitado_em, "
                           "sucesso) VALUES (?, 'busca_web', 'p', '{}', ?, ?, 1)", (a[3], "k" + quando, quando)).lastrowid
    def resultado(cid, titulo, data):
        con.execute("INSERT INTO resultados_web (chamada_id, run_id, consulta_id, uf, url_id, posicao, titulo, data_pub) "
                    "VALUES (?, ?, 'X', ?, ?, 1, ?, ?)", (cid, a[3], a[2], a[1], titulo, data))
    nova, velha = chamada("2099-01-02T00:00:00+00:00"), chamada("2098-01-01T00:00:00+00:00")
    resultado(nova, "mais recente, número menor", "2024-03-01")      # número menor, mas a chamada mais nova
    resultado(velha, "mais antigo, número maior", "2023-01-01")
    con.commit()
    l = next(x for x in LeituraPainelSQLite(banco).linhas_casos_estaduais() if x["achado_id"] == a[0])
    assert (l["titulo"], l["data_pub"]) == ("mais recente, número menor", "2024-03-01")
    resultado(nova, "empate de horário, número maior", "2024-04-01")    # mesmo horário: vale o número maior
    con.commit()
    l = next(x for x in LeituraPainelSQLite(banco).linhas_casos_estaduais() if x["achado_id"] == a[0])
    assert l["titulo"] == "empate de horário, número maior"


# ---------- painel e comando ----------
def _cli_painel(demo, tmp_path, autenticacao=False, contas=None):
    banco, raw = demo
    ctx = Contexto(banco="real", nome_banco="demo", somente_leitura=False, pasta_backup=tmp_path / "b", ufs=dict(C.UFS),
                   status_permitidos=set(V.STATUS), status_definicoes=dict(V.STATUS), hosts_permitidos=("testserver",),
                   autenticacao=autenticacao)
    return TestClient(criar_app(LeituraPainelSQLite(banco), lambda: RepositorioSQLite(banco), BrutoLocal(raw), ctx, contas=contas))


def test_download_pelo_painel_e_igual_ao_do_comando(demo, tmp_path, ambiente_cli, monkeypatch, capsys):
    banco, _ = demo
    with sqlite3.connect(banco) as c:          # um descarte desfeito volta a contar, no arquivo e no total impresso
        c.execute("UPDATE achados SET verificado = 0, status_final = 'descartado' WHERE achado_id = (SELECT MIN(achado_id) FROM v_achados_revisao)")
    r = _cli_painel(demo, tmp_path).get("/achados/casos.xlsx")
    assert r.status_code == 200 and r.headers["cache-control"] == "no-store"
    assert re.fullmatch(r'attachment; filename="casos_estaduais_\d{8}_\d{4}\.xlsx"', r.headers["content-disposition"])
    (ambiente_cli / "data").mkdir(parents=True, exist_ok=True)
    with sqlite3.connect(banco) as origem, sqlite3.connect(ambiente_cli / "data" / "simulado.db") as destino:
        origem.backup(destino)                  # leva tambem o que ainda esta no WAL
    monkeypatch.setattr(sys, "argv", ["main.py", "exportar-casos", "--simulado"])
    main.main()
    saida = capsys.readouterr().out
    arquivo = re.search(r"-> (.+\.xlsx)$", saida.strip()).group(1)
    assert "simulado_casos_estaduais_" in arquivo
    painel, comando = _abrir(r.content), openpyxl.load_workbook(arquivo)
    for aba in ("Casos Estaduais", "Resumo Casos"):
        assert [[c.value for c in l] for l in painel[aba].iter_rows()] == [[c.value for c in l] for l in comando[aba].iter_rows()]
    n = sqlite3.connect(banco).execute("SELECT COUNT(*) FROM v_achados_revisao WHERE NOT (verificado = 1 AND COALESCE(status_final, '') = 'descartado')").fetchone()[0]
    assert f"Casos estaduais: {n} achados ->" in saida and painel["Casos Estaduais"].max_row == n + 1


def test_botao_aparece_so_com_a_permissao_de_verificar(demo, tmp_path):
    assert 'href="/achados/casos.xlsx"' in _cli_painel(demo, tmp_path).get("/achados").text


def test_permissao_para_baixar():
    from dominio.autenticacao import permissao_exigida
    assert permissao_exigida("GET", "/achados/casos.xlsx") == "verificar_achados"
    assert permissao_exigida("GET", "/achados/12") == "ver_painel"


def test_sem_banco_mostra_a_pagina_de_sem_banco(tmp_path):
    ctx = Contexto(banco="real", nome_banco="x", somente_leitura=False, pasta_backup=tmp_path / "b", ufs=dict(C.UFS),
                   status_permitidos=set(V.STATUS), status_definicoes=dict(V.STATUS), hosts_permitidos=("testserver",))
    r = TestClient(criar_app(LeituraPainelSQLite(tmp_path / "nao.db"), lambda: None, BrutoLocal(tmp_path), ctx)).get("/achados/casos.xlsx")
    assert r.status_code == 503


def test_rotulos_do_painel_cobrem_o_vocabulario():
    assert set(rotulos.TECNOLOGIAS) == set(V.TECNOLOGIAS)
    assert set(V.STATUS) | set(V.STATUS_FINAL_EXTRA) <= set(rotulos.STATUS)


# ---------- permissões com login ----------
SENHA = "senha-de-teste-123"


def _com_login(demo, tmp_path, login):
    import re as _re
    from adaptadores.contas_sqlite import ContasSQLite
    from dominio import autenticacao as A
    contas = ContasSQLite(tmp_path / "contas.db")
    contas.garantir_padrao()
    grupos = {g["nome"]: g["grupo_id"] for g in contas.grupos()}
    agora = "2026-10-21T00:00:00+00:00"
    for nome, grupo in (("revisor", "Revisor"), ("leitor", "Leitor")):
        contas.criar_usuario(nome, nome, grupos[grupo], A.hash_senha(SENHA), agora)
    cli = _cli_painel(demo, tmp_path, autenticacao=True, contas=contas)
    if login:
        csrf = _re.search(r'name="csrf" value="([^"]+)"', cli.get("/entrar").text).group(1)
        assert cli.post("/entrar", data={"login": login, "senha": SENHA, "proximo": "/", "csrf": csrf},
                        follow_redirects=False).status_code == 303
    return cli


def test_revisor_baixa_e_ve_o_botao(demo, tmp_path):
    cli = _com_login(demo, tmp_path, "revisor")
    assert cli.get("/achados/casos.xlsx").status_code == 200
    assert 'href="/achados/casos.xlsx"' in cli.get("/achados").text


def test_leitor_nao_baixa_nem_ve_o_botao(demo, tmp_path):
    cli = _com_login(demo, tmp_path, "leitor")
    assert cli.get("/achados/casos.xlsx", follow_redirects=False).status_code == 403
    assert 'href="/achados/casos.xlsx"' not in cli.get("/achados").text


def test_visitante_anonimo_vai_para_a_entrada(demo, tmp_path):
    r = _com_login(demo, tmp_path, None).get("/achados/casos.xlsx", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/entrar?proximo=%2Fachados%2Fcasos.xlsx"
