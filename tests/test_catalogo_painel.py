"""Catalogo de referencias baixado pelo painel, no modelo do professor (uma aba por categoria, como as planilhas mestre dele)
ou no do projeto; o mesmo conteudo do comando exportar-referencias. Visitante nao baixa (o catalogo leva as notas de leitura)."""
import io
import shutil

import openpyxl
import pytest
from fastapi.testclient import TestClient

from adaptadores.bruto_local import BrutoLocal
from adaptadores.leitura_sqlite import LeituraPainelSQLite
from adaptadores.repositorio_sqlite import RepositorioSQLite
from aplicacao import catalogo, exportacao
from config import consultas as C, vocabulario as V
from config.categorias_referencias import CATEGORIAS_REFERENCIAS
from config.eixos_tematicos import EIXOS_TEMATICOS
from config.modelos_exportacao import MODELOS_REFERENCIAS
from painel.app import Contexto, criar_app
from test_autenticacao import _criar, _entrar, painel  # noqa: F401
from test_painel import demo  # noqa: F401

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
ABAS_PROFESSOR = ["Todas as Referencias", "Brasil", "Brasil_Casos", "Brasil_Etica", "Geral", "Inteligencia Policial",
                  "Internacional", "Internacional_Etica", "Legenda", "Indice_Tematico"]
ABAS_PADRAO = ["Referencias", "Referencias_Brasil", "Referencias_Internacional", "Legenda", "Indice_Tematico", "Colunas"]


def _ref(id_, categoria, pais="Brasil"):
    return {"id": id_, "categoria": categoria, "autor": "", "ano": 2024, "titulo": f"T{id_}", "tipo_fonte": "", "pais_escopo": pais,
            "topicos": "", "url": "", "prioridade": "", "status_leitura": "", "notas": "", "eixo": "", "citacoes": 0, "termos": "",
            "triagem_modelo": "", "decisao_humana": ""}


# ====================== abas do modelo do professor ======================
def test_modelo_do_professor_tem_uma_aba_por_categoria():
    cat = [_ref(1, "Geral", "Internacional"), _ref(2, "Brasil_Casos"), _ref(3, "Brasil_Casos"), _ref(4, ""), _ref(5, "brasil"),
           _ref(6, " Geral")]
    abas = exportacao.planilhas_referencias(cat, MODELOS_REFERENCIAS["professor"], CATEGORIAS_REFERENCIAS, EIXOS_TEMATICOS)
    assert list(abas) == ABAS_PROFESSOR
    assert abas["Todas as Referencias"]["ID"].tolist() == [1, 2, 3, 4, 5, 6]         # categoria vazia ou fora das 7: so em Todas
    assert abas["Brasil_Casos"]["ID"].tolist() == [2, 3]
    assert abas["Geral"]["ID"].tolist() == [1]
    assert abas["Brasil"]["ID"].tolist() == []                                       # pela Categoria, nao pelo pais
    assert all(list(abas[n].columns) == [c for c, _ in MODELOS_REFERENCIAS["professor"].colunas] for n in ABAS_PROFESSOR[:8])


def test_abas_do_professor_acompanham_a_lista_de_categorias():
    assert [n for n, _ in MODELOS_REFERENCIAS["professor"].abas[1:]] == sorted(CATEGORIAS_REFERENCIAS)


def test_catalogo_vazio_gera_abas_so_com_cabecalho():
    abas = exportacao.planilhas_referencias([], MODELOS_REFERENCIAS["professor"], CATEGORIAS_REFERENCIAS, EIXOS_TEMATICOS)
    wb = openpyxl.load_workbook(io.BytesIO(exportacao.planilha_xlsx(abas)), read_only=True)
    assert wb.sheetnames == ABAS_PROFESSOR and all(wb[n].max_row == 1 for n in ABAS_PROFESSOR[:8])


@pytest.mark.parametrize("titulo,gravado", [("=HYPERLINK(\"http://x\")", " =HYPERLINK(\"http://x\")"), ("a\x07b", "ab"),
                                            ("x" * 40000, "x" * (32000 - len("…[truncado]")) + "…[truncado]"),
                                            ("-1 e +2", "-1 e +2"), ("normal", "normal")])
def test_celulas_de_texto_seguem_a_regra_da_exportacao_principal(titulo, gravado):
    r = _ref(1, "Geral")
    r["titulo"] = titulo
    abas = exportacao.planilhas_referencias([r], MODELOS_REFERENCIAS["professor"], CATEGORIAS_REFERENCIAS, EIXOS_TEMATICOS)
    wb = openpyxl.load_workbook(io.BytesIO(exportacao.planilha_xlsx(abas)))
    celula = wb["Todas as Referencias"]["E2"]
    assert celula.value == gravado and celula.data_type == "s"


def test_reimportacao_le_so_a_primeira_aba(tmp_path):
    r = _ref(1, "Geral")
    abas = exportacao.planilhas_referencias([r], MODELOS_REFERENCIAS["professor"], CATEGORIAS_REFERENCIAS, EIXOS_TEMATICOS)
    abas["Geral"].loc[0, "Notas"] = "editado na aba da categoria"
    caminho = tmp_path / "c.xlsx"
    caminho.write_bytes(exportacao.planilha_xlsx(abas))
    linhas = exportacao.ler_leituras(caminho, MODELOS_REFERENCIAS["professor"], CATEGORIAS_REFERENCIAS, EIXOS_TEMATICOS)
    assert all("editado" not in str(l) for l in linhas)


def test_modelo_do_projeto_continua_como_era():
    cat = [_ref(1, "Geral", "Internacional"), _ref(2, "Brasil_Casos")]
    abas = exportacao.planilhas_referencias(cat, MODELOS_REFERENCIAS["padrao"], CATEGORIAS_REFERENCIAS, EIXOS_TEMATICOS)
    assert list(abas) == ABAS_PADRAO
    assert abas["Referencias_Brasil"]["ID"].tolist() == [2] and abas["Referencias_Internacional"]["ID"].tolist() == [1]


def test_planilha_em_memoria_tem_as_abas_na_ordem():
    abas = exportacao.planilhas_referencias([_ref(1, "Geral")], MODELOS_REFERENCIAS["professor"], CATEGORIAS_REFERENCIAS, EIXOS_TEMATICOS)
    wb = openpyxl.load_workbook(io.BytesIO(exportacao.planilha_xlsx(abas)), read_only=True)
    assert wb.sheetnames == ABAS_PROFESSOR


# ====================== download pelo painel ======================
class _FabricaProibida:
    def __call__(self):
        raise AssertionError("o download le pela leitura do painel, sem abrir o repositorio")


@pytest.fixture
def cli(demo, tmp_path):
    banco, raw = demo
    copia = tmp_path / "demo.db"
    shutil.copy(banco, copia)
    ctx = Contexto(banco="demo", nome_banco=copia.name, somente_leitura=False, pasta_backup=tmp_path / "backups", ufs=dict(C.UFS),
                   status_permitidos=set(V.STATUS) | set(V.STATUS_FINAL_EXTRA), status_definicoes=dict(V.STATUS),
                   hosts_permitidos=("testserver",), modelos_referencias=dict(MODELOS_REFERENCIAS),
                   categorias_referencias=dict(CATEGORIAS_REFERENCIAS), eixos_tematicos=dict(EIXOS_TEMATICOS))
    app = criar_app(LeituraPainelSQLite(copia), _FabricaProibida(), BrutoLocal(raw), ctx)
    return TestClient(app), copia


@pytest.mark.parametrize("modelo,arquivo,abas", [("professor", "referencias_modelo_professor.xlsx", ABAS_PROFESSOR),
                                                 ("padrao", "referencias.xlsx", ABAS_PADRAO)])
def test_download_traz_o_mesmo_catalogo_do_comando(cli, modelo, arquivo, abas):
    c, banco = cli
    r = c.get(f"/literatura/catalogo.xlsx?modelo={modelo}")
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    assert r.headers["content-disposition"] == f'attachment; filename="{arquivo}"'
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    assert wb.sheetnames == abas
    assert r.headers["cache-control"] == "no-store"
    esperado = catalogo.montar_catalogo(RepositorioSQLite(banco).linhas_catalogo())
    assert esperado and wb[abas[0]].max_row - 1 == len(esperado)
    if modelo == "professor":                                                  # cada aba de categoria pela rota
        for nome in abas[1:8]:
            ids = [row[0] for row in wb[nome].iter_rows(min_row=2, values_only=True)]
            assert ids == [x["id"] for x in esperado if x["categoria"] == nome], nome


@pytest.mark.parametrize("consulta", ["", "?modelo=", "?modelo=outro", "?modelo=PROFESSOR"])
def test_modelo_desconhecido_e_recusado(cli, consulta):
    c, _ = cli
    r = c.get("/literatura/catalogo.xlsx" + consulta)
    assert r.status_code == 404 and "Modelo de planilha desconhecido" in r.text
    assert "Use um dos modelos: padrao, professor." in r.text


def test_cartao_com_os_dois_botoes_na_visao_da_literatura(cli):
    c, _ = cli
    h = c.get("/literatura").text
    assert '<h2 id="catalogo-titulo" data-ic="download-simple" data-tom="rosa">Catálogo de referências</h2>' in h
    assert '<a class="botao" href="/literatura/catalogo.xlsx?modelo=professor">Baixar no modelo do professor</a>' in h
    assert '<a class="botao secundario" href="/literatura/catalogo.xlsx?modelo=padrao">Baixar no modelo do projeto</a>' in h


def test_sem_modelos_no_contexto_nao_ha_cartao(demo, tmp_path):
    from test_painel import _app
    app, _ = _app(demo[0], demo[1], tmp_path)
    assert "catalogo-titulo" not in TestClient(app).get("/literatura").text


def test_visitante_nao_baixa_e_leitor_recebe_403(painel):
    cli = painel[0]
    r = cli.get("/literatura/catalogo.xlsx?modelo=professor", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith("/entrar")
    _criar(painel, "leitora", "Leitor")
    assert _entrar(cli, "leitora").status_code == 303
    r = cli.get("/literatura/catalogo.xlsx?modelo=professor", follow_redirects=False)
    assert r.status_code == 403 and "Sem permissão" in r.text


def test_sem_banco_responde_como_as_outras_paginas(tmp_path):
    ctx = Contexto(banco="demo", nome_banco="x.db", somente_leitura=False, pasta_backup=tmp_path, ufs=dict(C.UFS),
                   status_permitidos=set(V.STATUS), status_definicoes=dict(V.STATUS), hosts_permitidos=("testserver",),
                   modelos_referencias=dict(MODELOS_REFERENCIAS))
    c = TestClient(criar_app(LeituraPainelSQLite(tmp_path / "nao_existe.db"), _FabricaProibida(), BrutoLocal(tmp_path), ctx))
    assert c.get("/literatura/catalogo.xlsx?modelo=professor").status_code == c.get("/literatura").status_code == 503


def test_revisor_baixa_e_leitor_nao_ve_o_cartao(painel):
    cli = painel[0]
    _criar(painel, "revisora", "Revisor")
    assert _entrar(cli, "revisora").status_code == 303
    r = cli.get("/literatura/catalogo.xlsx?modelo=professor")
    assert r.status_code == 404 and "Modelo de planilha desconhecido" in r.text   # passou da permissao; este contexto nao tem modelos
    cli.cookies.clear()
    _criar(painel, "leitora", "Leitor")
    assert _entrar(cli, "leitora").status_code == 303
    assert "catalogo-titulo" not in cli.get("/literatura").text


def test_uma_so_regra_de_celula_para_as_duas_exportacoes():
    from pathlib import Path
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from adaptadores import exportador_excel
    from dominio import planilha
    assert exportador_excel._limpar is planilha.celula_segura
    assert all(planilha.celula_segura(chr(i) + "x") == ILLEGAL_CHARACTERS_RE.sub("", chr(i) + "x") for i in range(0, 128) if chr(i) != "=")
    fonte = (Path(__file__).resolve().parent.parent / "aplicacao" / "exportacao.py").read_text(encoding="utf-8")
    assert "def _celula_segura" not in fonte and "LIMITE_CELULA =" not in fonte and "MARCADOR =" not in fonte
