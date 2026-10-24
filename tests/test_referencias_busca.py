"""Tela "Referências" (RB31, primeira etapa): as referências da busca de viés algorítmico nas colunas do modelo de
referências, com filtros e download; sem edição nesta etapa."""
import io
import logging
import re

import openpyxl
import pytest
from fastapi.testclient import TestClient

from adaptadores.bruto_local import BrutoLocal
from adaptadores.leitura_sqlite import LeituraPainelSQLite
from adaptadores.repositorio_sqlite import RepositorioSQLite
from aplicacao import busca_literatura, busca_academica
from aplicacao.comum import registrar_chamada
import config.settings as S
from config import consultas as C, vocabulario as V
from config.eixos_tematicos import EIXOS_EXPLORACAO
from dominio import referencias_busca as RB
from dominio.modelos import ArtigoAcademico, FiltrosAcademicos, FiltrosBusca, PaginaAcademica, RegistroChamada, RespostaBusca, ResultadoWeb
from infra.util import agora
from painel.app import Contexto, criar_app

CONJUNTO = {t.id: n for n, ts in C.CONJUNTOS_LITERATURA.items() for t in ts}


def _web(**kw):
    l = {"referencia_id": 1, "dominio": "exemplo.org", "url_original": "https://exemplo.org/a", "titulo": "T",
         "data_pub": "2024-05-01", "consultas": "LA2,LA1,LA2", "conjuntos": "A",
         "execucoes": "20261022_151331_c84423,20261021_213559_3526e4"}
    l.update(kw)
    return l


def _art(**kw):
    l = {"artigo_id": 7, "autores_publicacao": "M Silva, J Souza - Revista X, 2023 - scielo.br", "ano": 2023,
         "titulo": "Um artigo", "link": "https://scielo.br/x", "consultas": "LB1,LA3",
         "execucoes": "20261022_151407_ef2c8f"}
    l.update(kw)
    return l


# ---------- domínio ----------
def test_linha_da_web():
    assert RB.linha_web(_web()) == {
        "ID": "W-1", "Autor/Instituição": "exemplo.org", "Ano": 2024, "Título": "T", "Tipo de fonte": "", "País/escopo": "",
        "Tópicos e palavras-chave": "", "URL": "https://exemplo.org/a", "Prioridade de leitura": "",
        "Status de leitura": "Não iniciado", "Notas": "", "Fonte da busca": "Web", "Conjunto": "A", "Recorte": "", "Local": "",
        "Consultas": ["LA1", "LA2"], "Execução": ["20261021_213559_3526e4", "20261022_151331_c84423"]}


def test_linha_do_artigo():
    assert RB.linha_artigo(_art(), CONJUNTO) == {
        "ID": "A-7", "Autor/Instituição": "M Silva, J Souza", "Ano": 2023, "Título": "Um artigo", "Tipo de fonte": "",
        "País/escopo": "", "Tópicos e palavras-chave": "", "URL": "https://scielo.br/x",
        "Prioridade de leitura": "", "Status de leitura": "Não iniciado", "Notas": "", "Fonte da busca": "Scholar", "Conjunto": "A, B",
        "Recorte": "", "Local": "", "Consultas": ["LA3", "LB1"], "Execução": ["20261022_151407_ef2c8f"]}


@pytest.mark.parametrize("campo, valor, coluna, esperado", [
    ("dominio", None, "Autor/Instituição", ""), ("titulo", None, "Título", ""), ("url_original", None, "URL", ""),
    ("data_pub", None, "Ano", None), ("data_pub", "sem data", "Ano", None), ("consultas", None, "Tópicos e palavras-chave", ""), ("consultas", "LA1,LA2", "Tópicos e palavras-chave", ""),
    ("conjuntos", None, "Conjunto", ""), ("conjuntos", "B,A,B", "Conjunto", "A, B"),
    ("execucoes", None, "Execução", []), ("execucoes", "r2,r1,r2", "Execução", ["r1", "r2"]),
])
def test_linha_da_web_com_campos_vazios(campo, valor, coluna, esperado):
    assert RB.linha_web(_web(**{campo: valor}))[coluna] == esperado


@pytest.mark.parametrize("campo, valor, coluna, esperado", [
    ("autores_publicacao", None, "Autor/Instituição", ""), ("autores_publicacao", "Só autores", "Autor/Instituição", "Só autores"),
    ("ano", None, "Ano", None), ("ano", "2023", "Ano", None), ("consultas", "LX9", "Conjunto", ""),
])
def test_linha_do_artigo_com_campos_vazios(campo, valor, coluna, esperado):
    assert RB.linha_artigo(_art(**{campo: valor}), CONJUNTO)[coluna] == esperado


def test_ordem_do_ano_mais_recente_e_sem_ano_no_fim():
    web = [_web(referencia_id=1, data_pub="2020-01-01"), _web(referencia_id=2, data_pub=None), _web(referencia_id=10, data_pub="2024-01-01"),
           _web(referencia_id=3, data_pub="2024-02-02")]
    arts = [_art(artigo_id=5, ano=2024), _art(artigo_id=6, ano=None)]
    assert [r["ID"] for r in RB.referencias(web, arts, CONJUNTO)] == ["A-5", "W-3", "W-10", "W-1", "A-6", "W-2"]


@pytest.mark.parametrize("filtros, ids", [
    ({}, ["W-1", "W-2", "A-7"]),
    ({"conjunto": "A"}, ["W-1", "A-7"]), ({"conjunto": "B"}, ["W-2", "A-7"]),
    ({"fonte": "Scholar"}, ["A-7"]), ({"fonte": "Web"}, ["W-1", "W-2"]),
    ({"texto": "  AVALIAÇÃO  "}, ["W-2"]), ({"texto": "avaliacao"}, ["W-2"]), ({"texto": "silva"}, ["A-7"]),
    ({"texto": "scielo.br/x"}, ["A-7"]), ({"conjunto": "B", "fonte": "Web"}, ["W-2"]), ({"texto": "nada"}, []),
    ({"texto": "   "}, ["W-1", "W-2", "A-7"]),                    # só espaços: sem filtro
    ({"texto": "ø"}, []), ({"texto": "中"}, []),                   # a chave sem acento descarta tudo: nada casa, não tudo
])
def test_filtros(filtros, ids):
    linhas = [RB.linha_web(_web(referencia_id=1, titulo="Uso de IA")), RB.linha_web(_web(referencia_id=2, titulo="Avaliação", conjuntos="B")),
              RB.linha_artigo(_art(), CONJUNTO)]
    assert [r["ID"] for r in RB.filtrar(linhas, **filtros)] == ids


# ---------- banco ----------
class _Busca:
    nome = "falsa"

    def __init__(self, resultados):
        self._r = resultados

    def buscar(self, consulta, filtros):
        return RespostaBusca(self._r(consulta), RegistroChamada(provedor="falsa", operacao="buscar", parametros={},
                                                                 requisitado_em=agora(), http_status=200, raw=b"{}"))


class _Scholar:
    nome = "falso"

    def buscar_pagina(self, termo, inicio, filtros):
        arts = [ArtigoAcademico(f"Artigo {termo.id}", "https://s.org/" + termo.id, "A Autor - Rev, 2022", 2022, "r", 0, 1)]
        return PaginaAcademica(arts, 1, RegistroChamada(provedor="falso", operacao="buscar_pagina", parametros={},
                                                         requisitado_em=agora(), http_status=200, raw=b"{}"))


@pytest.fixture
def banco(tmp_path):
    logging.disable(logging.CRITICAL)
    caminho = tmp_path / "busca.db"
    repo = RepositorioSQLite(caminho)
    try:
        snap = S.snapshot_config("vies")                # a configuração de verdade, para a página da execução
        run = repo.iniciar_execucao("buscar-literatura", snap["hash"], snap["json"], "c")
        busca_literatura.executar(repo, BrutoLocal(tmp_path / "raw"), _Busca(lambda c: [
            ResultadoWeb("https://comum.org/x", f"comum por {c.id}", "s", "2025-01-01", 1),
            ResultadoWeb(f"https://so.org/{c.id}", f"só {c.id}", "s", None, 2)]), run,
            [("A", C.CONJUNTOS_LITERATURA["A"][:2]), ("B", C.CONJUNTOS_LITERATURA["B"][:1])], FiltrosBusca(max_resultados=20), [], 0)
        busca_academica.executar(repo, BrutoLocal(tmp_path / "raw"), _Scholar(), run, C.CONJUNTOS_LITERATURA["B"][:1],
                                 FiltrosAcademicos(por_pagina=20), 1, [], 0)
        busca_academica.executar(repo, BrutoLocal(tmp_path / "raw"), _Scholar(), run, C.TERMOS[:1],
                                 FiltrosAcademicos(por_pagina=20), 1, [], 0)      # um termo da Tarefa 2: não entra
        repo.finalizar_execucao(run, "concluida")
    finally:
        repo.fechar()
        logging.disable(logging.NOTSET)
    return caminho


def test_leitura_das_referencias_web(banco):
    linhas = {l["url_original"]: l for l in LeituraPainelSQLite(banco).linhas_referencias_web()}
    comum = linhas["https://comum.org/x"]
    assert sorted(comum["consultas"].split(",")) == ["LA1", "LA2", "LB1"] and sorted(comum["conjuntos"].split(",")) == ["A", "B"]
    assert comum["titulo"] == "comum por LB1"                      # o resultado da chamada mais recente
    assert linhas["https://so.org/LA2"]["consultas"] == "LA2" and len(linhas) == 4


def test_referencia_sem_ocorrencia_aparece_sem_consulta_nem_conjunto(banco):
    import sqlite3
    with sqlite3.connect(banco) as c:
        c.execute("INSERT INTO referencias_web (url_original, url_normalizada, dominio, primeira_vez_run, primeira_vez_em) VALUES ('https://orfa.org/', 'orfa.org', 'orfa.org', 'r', '2026-01-01')")
    linha = [l for l in LeituraPainelSQLite(banco).linhas_referencias_web() if l["url_original"] == "https://orfa.org/"][0]
    assert (linha["titulo"], linha["consultas"], linha["conjuntos"]) == (None, None, None)
    r = RB.linha_web(linha)
    assert (r["Título"], r["Conjunto"], r["Tópicos e palavras-chave"], r["Ano"]) == ("", "", "", None)


def test_artigos_das_consultas_numa_passada_so(banco):
    """Sem subconsulta correlacionada por artigo (ocorrencias não tem índice por artigo_id)."""
    import sqlite3
    from adaptadores.leitura_sqlite import SQL_ARTIGOS_DAS_CONSULTAS
    plano = sqlite3.connect(banco).execute("EXPLAIN QUERY PLAN " + SQL_ARTIGOS_DAS_CONSULTAS.format(marcas="?, ?"), ("LA1", "LB1")).fetchall()
    assert plano and not any("CORRELATED" in str(linha[-1]) for linha in plano)


def test_leitura_dos_artigos_so_das_consultas_da_busca(banco):
    leitura = LeituraPainelSQLite(banco)
    assert [(l["titulo"], l["consultas"]) for l in leitura.linhas_artigos_das_consultas(list(CONJUNTO))] == [("Artigo LB1", "LB1")]
    assert leitura.linhas_artigos_das_consultas([]) == []


# ---------- painel ----------
def _cli(banco, tmp_path):
    ctx = Contexto(banco="vies", nome_banco="busca.db", somente_leitura=True, pasta_backup=tmp_path / "b", ufs=dict(C.UFS),
                   status_permitidos=set(V.STATUS), status_definicoes=dict(V.STATUS), hosts_permitidos=("testserver",),
                   conjunto_da_consulta=dict(CONJUNTO), eixos_exploracao=dict(EIXOS_EXPLORACAO))
    return TestClient(criar_app(LeituraPainelSQLite(banco), lambda: RepositorioSQLite(banco), BrutoLocal(tmp_path), ctx))


def _lista(html):
    """A pagina sem a legenda do fim (que tem a sua propria tabela de consultas)."""
    return html.split('<section class="legenda-pagina"')[0]


def test_tela_mostra_as_colunas_do_modelo_e_conta(banco, tmp_path):
    r = _cli(banco, tmp_path).get("/literatura/referencias")
    assert r.status_code == 200
    cabecalhos = re.findall(r"<th[^>]*>([^<]+)</th>", _lista(r.text))
    # as 11 colunas do modelo, a execução logo ao lado delas, depois a fonte da busca e o conjunto
    assert cabecalhos == list(RB.COLUNAS) + ["Recorte", "Local", "Consultas", "Execução", "Fonte da busca", "Conjunto"]
    assert "<strong>5</strong> referência(s)" in r.text
    assert 'href="/literatura/referencias"' in _cli(banco, tmp_path).get("/").text          # item do menu


@pytest.mark.parametrize("consulta, n", [("conjunto=B", 3), ("fonte=Scholar", 1), ("q=s%C3%B3", 3), ("conjunto=A&fonte=Scholar", 0)])
def test_filtros_na_tela(consulta, n, banco, tmp_path):
    r = _cli(banco, tmp_path).get(f"/literatura/referencias?{consulta}")
    assert r.status_code == 200 and f"<strong>{n}</strong> referência(s)" in r.text


@pytest.mark.parametrize("consulta", ["conjunto=D", "fonte=Google", "p=0", "p=abc", "conjunto=+A", "conjunto=a", "fonte=web", "fonte=Perplexity",
                                      "fonte=Scholar+"])
def test_filtro_invalido_da_404(consulta, banco, tmp_path):
    assert _cli(banco, tmp_path).get(f"/literatura/referencias?{consulta}").status_code == 404


def test_paginacao_de_50(tmp_path):
    caminho = tmp_path / "muitas.db"
    repo = RepositorioSQLite(caminho)
    try:
        run = repo.iniciar_execucao("buscar-literatura", "h", "{}", "c")
        reg = RegistroChamada(provedor="f", operacao="buscar", parametros={}, requisitado_em=agora(), http_status=200)
        cid, _ = registrar_chamada(repo, None, run, "busca_literatura", "consulta_literatura", "LA1", "k", reg, [])
        repo.salvar_referencias_web(cid, run, "A", "LA1", [ResultadoWeb(f"https://e.org/{i}", f"t{i}", "", None, i) for i in range(1, 61)])
        repo.confirmar()
    finally:
        repo.fechar()
    cli = _cli(caminho, tmp_path)
    assert _lista(cli.get("/literatura/referencias").text).count("<tr>") == 51          # cabeçalho + 50
    p2 = cli.get("/literatura/referencias?p=2").text
    assert _lista(p2).count("<tr>") == 11 and "Página 2 de 2" in p2
    # sem ano, empate pelo número: a página 2 tem as referências 51 a 60, nessa ordem
    assert re.findall(r'<a href="/literatura/referencias/web/\d+">(W-\d+)</a>', p2) == [f"W-{i}" for i in range(51, 61)]
    p3 = cli.get("/literatura/referencias?p=3")                      # além da última: como nas outras listas, vazia
    assert p3.status_code == 200 and _lista(p3.text).count("<tr>") == 1 and "Página 3 de 2" in p3.text
    filtrada = cli.get("/literatura/referencias?fonte=Web&q=t").text
    assert 'href="/literatura/referencias?fonte=Web&amp;q=t&amp;p=2"' in filtrada
    assert 'href="/literatura/referencias.xlsx?fonte=Web&amp;q=t&amp;modelo=erico"' in filtrada


def test_referencias_clicaveis(banco, tmp_path):
    """O ID abre a página do item no painel (artigo ou referência da web); o título abre a fonte."""
    html = _cli(banco, tmp_path).get("/literatura/referencias").text
    artigo = re.search(r'<td><a href="/literatura/artigos/(\d+)">A-(\d+)</a></td>', html)
    assert artigo and artigo.group(1) == artigo.group(2)
    assert _cli(banco, tmp_path).get(f"/literatura/artigos/{artigo.group(1)}").status_code == 200
    assert '<td><a href="https://s.org/LB1" rel="noopener noreferrer" target="_blank">Artigo LB1</a></td>' in html
    assert '<td><a href="https://comum.org/x" rel="noopener noreferrer" target="_blank">comum por LB1</a></td>' in html
    web = re.search(r'<td><a href="/literatura/referencias/web/(\d+)">W-(\d+)</a></td>', html)     # a página da referência
    assert web and web.group(1) == web.group(2)


def test_leitura_traz_as_execucoes(banco):
    leitura = LeituraPainelSQLite(banco)
    run = leitura.linhas_referencias_web()[0]["execucoes"]
    assert run and "," not in run                                         # uma execução só na base de teste
    assert [l["execucoes"] for l in leitura.linhas_artigos_das_consultas(list(CONJUNTO))] == [run]


def test_execucao_ao_lado_das_colunas_do_modelo_e_clicavel(banco, tmp_path):
    cli = _cli(banco, tmp_path)
    html = cli.get("/literatura/referencias").text
    run = LeituraPainelSQLite(banco).linhas_referencias_web()[0]["execucoes"]
    link = f'<a href="/execucoes/{run}?configuracao=1#configuracao"><code>{run}</code></a>'
    assert html.count(link) == 5                                          # as 4 referências da web e o artigo
    pagina = cli.get(f"/execucoes/{run}?configuracao=1")
    assert pagina.status_code == 200 and '<details class="cartao" id="configuracao" open>' in pagina.text
    assert '<details class="cartao" id="configuracao">' in cli.get(f"/execucoes/{run}").text     # sem o parâmetro, fechada


def test_pagina_da_execucao_mostra_todos_os_parametros_da_busca(banco, tmp_path):
    import html as H
    run = LeituraPainelSQLite(banco).linhas_referencias_web()[0]["execucoes"]
    texto = H.unescape(_cli(banco, tmp_path).get(f"/execucoes/{run}?configuracao=1").text)
    assert "Banco: <strong>vies</strong>" in texto
    assert "Consultas da busca de literatura (por conjunto)" in texto and "Filtros da busca de literatura na web" in texto
    for conjunto, termos in C.CONJUNTOS_LITERATURA.items():
        for t in termos:
            assert f"<tr><td>{conjunto}</td><td>{t.id}</td><td>{t.query}</td><td>{t.idioma}</td></tr>" in texto
    assert "Prompt de sistema" in texto and "Prompt do usuário" in texto and "Provedores" in texto


def test_sem_permissao_de_ver_execucoes_a_execucao_e_so_texto(banco, tmp_path):
    from adaptadores.contas_sqlite import ContasSQLite
    from dominio import autenticacao as A
    contas = ContasSQLite(tmp_path / "contas.db")
    contas.garantir_padrao()
    grupos = {g["nome"]: g["grupo_id"] for g in contas.grupos()}
    contas.criar_usuario("leitor", "leitor", grupos["Leitor"], A.hash_senha("senha-de-teste-123"), "2026-10-22T00:00:00+00:00")
    ctx = Contexto(banco="vies", nome_banco="busca.db", somente_leitura=False, pasta_backup=tmp_path / "b", ufs=dict(C.UFS),
                   status_permitidos=set(V.STATUS), status_definicoes=dict(V.STATUS), hosts_permitidos=("testserver",),
                   conjunto_da_consulta=dict(CONJUNTO), eixos_exploracao=dict(EIXOS_EXPLORACAO), autenticacao=True)
    cli = TestClient(criar_app(LeituraPainelSQLite(banco), lambda: RepositorioSQLite(banco), BrutoLocal(tmp_path), ctx, contas=contas))
    csrf = re.search(r'name="csrf" value="([^"]+)"', cli.get("/entrar").text).group(1)
    assert cli.post("/entrar", data={"login": "leitor", "senha": "senha-de-teste-123", "proximo": "/", "csrf": csrf},
                    follow_redirects=False).status_code == 303
    html = cli.get("/literatura/referencias").text
    run = LeituraPainelSQLite(banco).linhas_referencias_web()[0]["execucoes"]
    assert f"<code>{run}</code>" in html and f'href="/execucoes/{run}' not in html
    assert cli.get(f"/execucoes/{run}", follow_redirects=False).status_code == 403


def test_texto_so_com_espacos_nao_vira_filtro(banco, tmp_path):
    r = _cli(banco, tmp_path).get("/literatura/referencias?q=+++")
    assert "<strong>5</strong> referência(s)" in r.text and "Limpar" not in r.text and 'href="/literatura/referencias.xlsx?modelo=erico"' in r.text


@pytest.mark.parametrize("url, esperado", [
    ("javascript:alert(1)", "endereço não é http/https"), ("data:text/html,x", "endereço não é http/https"),
    ("JAVASCRIPT:alert(1)", "endereço não é http/https"), ("https://ok.org/<b>", 'href="https://ok.org/&lt;b&gt;"'),
])
def test_url_que_nao_e_http_nao_vira_link(url, esperado, tmp_path):
    caminho = tmp_path / "xss.db"
    repo = RepositorioSQLite(caminho)
    try:
        run = repo.iniciar_execucao("buscar-literatura", "h", "{}", "c")
        reg = RegistroChamada(provedor="f", operacao="buscar", parametros={}, requisitado_em=agora(), http_status=200)
        cid, _ = registrar_chamada(repo, None, run, "busca_literatura", "consulta_literatura", "LA1", "k", reg, [])
        repo.salvar_referencias_web(cid, run, "A", "LA1", [ResultadoWeb(url, "<script>x</script>", "", None, 1)])
        repo.confirmar()
    finally:
        repo.fechar()
    html = _cli(caminho, tmp_path).get("/literatura/referencias").text
    assert esperado in html and "<script>x</script>" not in html and 'href="javascript' not in html.lower()
    assert 'href="data:' not in html.lower()


def test_download_tem_as_linhas_da_tela_na_mesma_ordem(banco, tmp_path):
    cli = _cli(banco, tmp_path)
    tela = re.findall(r">([WA]-\d+)</", cli.get("/literatura/referencias").text)
    wb = openpyxl.load_workbook(io.BytesIO(cli.get("/literatura/referencias.xlsx").content))
    assert [l[0].value for l in wb["Referências ampliadas"].iter_rows(min_row=2)] == tela and len(tela) == 5 and any(i.startswith("A-") for i in tela)


def test_nome_do_arquivo_no_banco_real_nao_tem_prefixo(banco, tmp_path):
    ctx = Contexto(banco="real", nome_banco="busca.db", somente_leitura=True, pasta_backup=tmp_path / "b", ufs=dict(C.UFS),
                   status_permitidos=set(V.STATUS), status_definicoes=dict(V.STATUS), hosts_permitidos=("testserver",),
                   conjunto_da_consulta=dict(CONJUNTO), eixos_exploracao=dict(EIXOS_EXPLORACAO))
    r = TestClient(criar_app(LeituraPainelSQLite(banco), lambda: None, BrutoLocal(tmp_path), ctx)).get("/literatura/referencias.xlsx")
    assert re.fullmatch(r'attachment; filename="referencias_\d{8}_\d{4}\.xlsx"', r.headers["content-disposition"])


def test_permissao_e_a_de_ver_o_painel():
    """Sem notas nem nomes de revisores nesta etapa: quem vê o painel vê a lista e baixa."""
    from dominio.autenticacao import permissao_exigida
    assert permissao_exigida("GET", "/literatura/referencias") == permissao_exigida("GET", "/literatura/referencias.xlsx") == "ver_painel"


def test_sem_banco_mostra_a_pagina_de_sem_banco(tmp_path):
    for rota in ("/literatura/referencias", "/literatura/referencias.xlsx"):
        assert _cli(tmp_path / "nao.db", tmp_path).get(rota).status_code == 503


def test_postgres_igual_ao_sqlite(banco, tmp_path, pg_url):
    from adaptadores import migrador_postgres as M
    from adaptadores.leitura_postgres import LeituraPainelPostgres
    from adaptadores.postgres import importador as I
    from adaptadores.postgres.conexao import chaves_do_banco, criar_pool
    M.aplicar(pg_url)
    I.importar(banco, pg_url)
    pool = criar_pool(pg_url, maximo=2)
    try:
        sq, pg = LeituraPainelSQLite(banco), LeituraPainelPostgres(pool, chaves_do_banco(pool))
        assert [dict(l) for l in pg.linhas_referencias_web()] == [dict(l) for l in sq.linhas_referencias_web()]
        assert [dict(l) for l in pg.linhas_artigos_das_consultas(list(CONJUNTO))] == \
            [dict(l) for l in sq.linhas_artigos_das_consultas(list(CONJUNTO))]
    finally:
        pool.close()


def test_download_tem_as_mesmas_linhas_e_o_esquema(banco, tmp_path):
    r = _cli(banco, tmp_path).get("/literatura/referencias.xlsx?conjunto=B")
    assert r.status_code == 200 and re.fullmatch(r'attachment; filename="vies_referencias_\d{8}_\d{4}\.xlsx"', r.headers["content-disposition"])
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    # o formato do modelo ampliado (ia_seguranca_publica_vies_algoritmico_investigacao_criminal_ampliado.xlsx)
    assert wb.sheetnames == ["Referências ampliadas", "Esquema de exploração"]
    ref, esq = wb["Referências ampliadas"], wb["Esquema de exploração"]
    assert [c.value for c in ref[1]] == list(RB.COLUNAS) and ref.max_row == 4 and ref.max_column == 11   # sem Execução, Fonte, Conjunto
    assert [ref.column_dimensions[l].width for l in "ABCDEFGHIJK"] == [12, 56, 12, 60, 59, 19, 60, 60, 23, 19, 12]
    assert [[c.value for c in l] for l in esq.iter_rows(min_row=2)] == [list(x) for x in EIXOS_EXPLORACAO.items()]
    assert [esq.column_dimensions[l].width for l in "AB"] == [35, 60] and (ref.freeze_panes, esq.freeze_panes) == ("A2", "A2")
    assert (ref.auto_filter.ref, esq.auto_filter.ref) == ("A1:K4", "A1:B7")
    for c in list(ref[1]) + list(esq[1]):            # cabeçalho: negrito, centralizado no topo, borda fina, sem fundo
        b = c.border
        assert (c.font.b, c.alignment.horizontal, c.alignment.vertical, c.fill.fill_type) == (True, "center", "top", None)
        assert (b.left.style, b.right.style, b.top.style, b.bottom.style) == ("thin",) * 4
    for linha in ref.iter_rows(min_row=2):
        assert linha[2].value is None or isinstance(linha[2].value, str)       # o ano é texto, como no modelo
        assert linha[9].value == "Não iniciado" and not linha[0].font.b


def test_eixos_do_esquema_sao_os_seis_literais():
    assert list(EIXOS_EXPLORACAO) == ["Policiamento preditivo e equidade", "Biometria e vigilância", "Investigação e inteligência",
                                      "Forense digital e deepfakes", "Direitos e governança", "Brasil e América Latina"]
    assert EIXOS_EXPLORACAO["Brasil e América Latina"] == ("Mapear pesquisas por UF, contratos e políticas; comparar dados de "
                                                          "implementação, resultados e impactos distributivos.")


def test_botao_baixar_curto_no_cartao_dos_filtros(banco, tmp_path):
    html = _cli(banco, tmp_path).get("/literatura/referencias?conjunto=B").text
    form = re.search(r'<form class="filtros".*?</form>', html, re.S).group(0)
    assert 'href="/literatura/referencias.xlsx?conjunto=B&amp;modelo=erico">Modelo Érico</a>' in form and ">Baixar</summary>" in form
    assert "Baixar no modelo de referências" not in html and "<strong>3</strong> referência(s).</p>" in html
