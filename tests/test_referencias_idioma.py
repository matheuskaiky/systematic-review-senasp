"""Tela Referências: a coluna e o filtro "Idioma da busca" dos artigos do Scholar (só português, sem filtro, as duas ou
outro filtro), pelas chamadas que acharam cada artigo; vazia nas referências da web; igual no SQLite e no PostgreSQL."""
import io
import logging
import re

import openpyxl
import pytest

from adaptadores.bruto_local import BrutoLocal
from adaptadores.leitura_sqlite import LeituraPainelSQLite
from adaptadores.repositorio_sqlite import RepositorioSQLite
from aplicacao import busca_academica, busca_literatura
from config import consultas as C
from dominio import referencias_busca as RB
from dominio.modelos import ArtigoAcademico, FiltrosAcademicos, FiltrosBusca, PaginaAcademica, RegistroChamada, ResultadoWeb
from infra.util import agora
from test_referencias_recorte import CONJUNTO, _cabecalhos, _cli
from test_triagem_vies import _Busca

A = {t.id: t for t in C.CONJUNTOS_LITERATURA["A"]}


@pytest.mark.parametrize("idiomas, rotulo", [
    (None, ""), ("", ""), (",", ""), (" , ", ""),
    ("so_portugues", "Só português"), ("so_portugues,so_portugues", "Só português"), (" so_portugues ", "Só português"),
    ("sem_filtro", "Sem filtro"),
    ("so_portugues,sem_filtro", "As duas"), ("sem_filtro,so_portugues", "As duas"),
    ("outro", "Outro filtro"), ("so_portugues,outro", "Outro filtro"), ("so_portugues,sem_filtro,outro", "Outro filtro"),
    ("SO_PORTUGUES", "Outro filtro"), ("lang_pt", "Outro filtro"),
    ("sem_filtro,sem_filtro", "Sem filtro"), ("outro,outro", "Outro filtro"), ("sem_filtro,xyz", "Outro filtro"),
])
def test_rotulo_pelas_chaves_das_chamadas(idiomas, rotulo):
    assert RB.idioma_da_busca(idiomas) == rotulo


def test_coluna_entre_a_fonte_e_o_conjunto_e_so_no_scholar():
    assert RB.COLUNAS_TELA[-3:] == ("Fonte da busca", "Idioma da busca", "Conjunto") and len(RB.COLUNAS_TELA) == 18
    assert RB.IDIOMAS_DA_BUSCA == {"so_portugues": "Só português", "sem_filtro": "Sem filtro", "as_duas": "As duas"}
    assert RB.LARGURA_DA_COLUNA["Idioma da busca"] == 16
    web = RB.linha_web({"referencia_id": 1, "dominio": "e.org", "data_pub": None, "titulo": "T", "url_original": "https://e.org",
                        "conjuntos": "A", "execucoes": "r", "consultas": "LA1"})
    art = RB.linha_artigo({"artigo_id": 1, "autores_publicacao": "X - Y", "ano": 2024, "titulo": "T", "link": "https://s.org",
                           "consultas": "LA1", "execucoes": "r", "idiomas": "sem_filtro,so_portugues"}, CONJUNTO)
    sem_chave = RB.linha_artigo({"artigo_id": 2, "autores_publicacao": "", "ano": None, "titulo": "", "link": "",
                                 "consultas": "LA1", "execucoes": "r"}, CONJUNTO)
    assert (web["Idioma da busca"], art["Idioma da busca"], sem_chave["Idioma da busca"]) == ("", "As duas", "")


class _ScholarComIdioma:
    """Um artigo por consulta (https://s.org/<id>); a chamada guarda o lr do filtro, como o adaptador real."""
    nome = "falso"

    def buscar_pagina(self, termo, inicio, filtros):
        params = {"q": termo.query} | ({"lr": filtros.idioma_resultados} if filtros.idioma_resultados else {})
        arts = [ArtigoAcademico(f"Artigo {termo.id}", "https://s.org/" + termo.id, "A Autor - Rev, 2022", 2022, "r", 0, 1)]
        return PaginaAcademica(arts, 1, RegistroChamada(provedor="falso", operacao="buscar_pagina", parametros=params,
                                                         requisitado_em=agora(), http_status=200, raw=b"{}"))


def _scholar(repo, raw, termos, idioma):
    run = repo.iniciar_execucao("academico", "h", "{}", "c")
    busca_academica.executar(repo, BrutoLocal(raw), _ScholarComIdioma(), run, [A[t] for t in termos],
                             FiltrosAcademicos(por_pagina=20, idioma_resultados=idioma), 1, [], 0)
    repo.finalizar_execucao(run, "concluida")


@pytest.fixture
def banco(tmp_path):
    """LA1 só com lang_pt; LA2 com lang_pt e sem filtro; LA3 só sem filtro; LA4 com lang_en; uma referência da web."""
    caminho, raw = tmp_path / "busca.db", tmp_path / "raw"
    logging.disable(logging.CRITICAL)
    repo = RepositorioSQLite(caminho)
    try:
        run = repo.iniciar_execucao("buscar-literatura", "h", "{}", "c")
        busca_literatura.executar(repo, BrutoLocal(raw), _Busca(lambda c: [ResultadoWeb("https://web.org/x", "web", "t", None, 1)]),
                                  run, [("A", C.CONJUNTOS_LITERATURA["A"][:1])], FiltrosBusca(max_resultados=20), [], 0)
        repo.finalizar_execucao(run, "concluida")
        _scholar(repo, raw, ["LA1", "LA2"], "lang_pt")
        _scholar(repo, raw, ["LA2", "LA3"], None)
        _scholar(repo, raw, ["LA4"], "lang_en")
    finally:
        repo.fechar()
        logging.disable(logging.NOTSET)
    return caminho


ESPERADO = {"https://s.org/LA1": "Só português", "https://s.org/LA2": "As duas", "https://s.org/LA3": "Sem filtro",
            "https://s.org/LA4": "Outro filtro", "https://web.org/x": ""}


def _idioma_por_url(html):
    linhas = re.findall(r"<tr>(.*?)</tr>", html.split('<section class="legenda-pagina"')[0], re.S)[1:]
    cab = _cabecalhos(html)
    saida = {}
    for l in linhas:
        celulas = re.findall(r"<td[^>]*>(.*?)</td>", l, re.S)
        url = re.search(r'href="(https://[^"]+)"', celulas[cab.index("URL")]).group(1)
        saida[url] = celulas[cab.index("Idioma da busca")]
    return saida


def test_tela_mostra_o_idioma_de_cada_artigo(banco, tmp_path):
    html = _cli(banco, tmp_path).get("/literatura/referencias").text
    assert _idioma_por_url(html) == ESPERADO


@pytest.mark.parametrize("idioma, urls", [
    ("so_portugues", ["https://s.org/LA1"]), ("sem_filtro", ["https://s.org/LA3"]), ("as_duas", ["https://s.org/LA2"]),
])
def test_filtro_por_idioma_deixa_de_fora_a_web_e_os_outros(banco, tmp_path, idioma, urls):
    html = _cli(banco, tmp_path).get(f"/literatura/referencias?idioma={idioma}").text
    assert sorted(_idioma_por_url(html)) == urls and "<strong>1</strong> referência(s)" in html
    assert f'<option value="{idioma}" selected>' in html


@pytest.mark.parametrize("idioma", ["outro", "Só português", "SO_PORTUGUES", " so_portugues", "lang_pt", "todos"])
def test_idioma_invalido_e_404(banco, tmp_path, idioma):
    r = _cli(banco, tmp_path).get("/literatura/referencias", params={"idioma": idioma})
    assert r.status_code == 404 and "Um dos filtros da página não é um valor aceito." in r.text
    assert _cli(banco, tmp_path).get("/literatura/referencias.xlsx", params={"idioma": idioma}).status_code == 404


def test_select_do_filtro_na_ordem(banco, tmp_path):
    form = re.search(r'<form class="filtros".*?</form>', _cli(banco, tmp_path).get("/literatura/referencias").text, re.S).group(0)
    assert ('<label>Idioma da busca <select name="idioma"><option value="">todos</option><option value="so_portugues">Só português'
            '</option><option value="sem_filtro">Sem filtro</option><option value="as_duas">As duas</option></select></label>') in form
    assert form.index('name="recorte"') < form.index('name="idioma"') < form.index('name="q"')


def test_endereco_repete_o_filtro_na_ordem_e_o_download_o_respeita(banco, tmp_path):
    cli = _cli(banco, tmp_path)
    html = cli.get("/literatura/referencias?q=Artigo&idioma=as_duas&recorte=vies_algoritmico&fonte=Scholar&conjunto=A").text
    assert ('href="/literatura/referencias.xlsx?conjunto=A&amp;fonte=Scholar&amp;recorte=vies_algoritmico&amp;idioma=as_duas'
            '&amp;q=Artigo&amp;modelo=colunas"') in html
    wb = openpyxl.load_workbook(io.BytesIO(cli.get("/literatura/referencias.xlsx?idioma=as_duas&modelo=colunas").content))
    folha = wb.active
    cab = [c.value for c in folha[1]]
    assert cab == list(RB.COLUNAS_TELA)
    linhas = [[c.value for c in r] for r in folha.iter_rows(min_row=2)]
    assert [l[cab.index("Idioma da busca")] for l in linhas] == ["As duas"] and [l[cab.index("URL")] for l in linhas] == ["https://s.org/LA2"]
    erico = openpyxl.load_workbook(io.BytesIO(cli.get("/literatura/referencias.xlsx?idioma=so_portugues").content)).active
    assert "Idioma da busca" not in [c.value for c in erico[1]] and erico.max_row == 2      # o modelo Érico só filtra


def test_postgres_igual_ao_sqlite(banco, pg_url):
    from adaptadores import migrador_postgres as M
    from adaptadores.leitura_postgres import LeituraPainelPostgres
    from adaptadores.postgres import importador as I
    from adaptadores.postgres.conexao import chaves_do_banco, criar_pool
    M.aplicar(pg_url)
    I.importar(banco, pg_url)
    pool = criar_pool(pg_url, maximo=2)
    try:
        sq = LeituraPainelSQLite(banco).linhas_artigos_das_consultas(list(CONJUNTO))
        pg = LeituraPainelPostgres(pool, chaves_do_banco(pool)).linhas_artigos_das_consultas(list(CONJUNTO))
        def _sem_ordem(linhas):          # a ordem do GROUP_CONCAT muda entre os bancos (e entre execucoes do mesmo segundo)
            return [{**dict(l), **{c: sorted((l[c] or "").split(",")) for c in ("idiomas", "execucoes", "consultas")}} for l in linhas]
        assert _sem_ordem(pg) == _sem_ordem(sq)
        assert {l["link"]: RB.idioma_da_busca(l["idiomas"]) for l in pg} == {u: r for u, r in ESPERADO.items() if "s.org" in u}
    finally:
        pool.close()


def _idioma_de(banco, link):
    linhas = LeituraPainelSQLite(banco).linhas_artigos_das_consultas(list(CONJUNTO))
    return next(l["idiomas"] for l in linhas if l["link"] == link)


@pytest.mark.parametrize("parametros, chave", [
    ('{"q": "x"}', "sem_filtro"), ('{"lr": null}', "sem_filtro"), ('{"lr": "lang_pt"}', "so_portugues"),
    ('{"lr": ""}', "outro"), ('{"lr": 1}', "outro"), ('{"lr": []}', "outro"), ('{"lr": {}}', "outro"),
    ('{"lr": "LANG_PT"}', "outro"), ('{"lr": " lang_pt"}', "outro"), ('[]', "outro"), ('"lang_pt"', "outro"),
    ('null', "outro"), ('xx', "outro"), ('', "outro"), ('{"lr": "a", "lr": "lang_pt"}', "so_portugues"),
])
def test_chave_de_cada_chamada_pelos_parametros_gravados(banco, parametros, chave):
    """O adaptador le os parametros da chamada em Python (o mesmo resultado no SQLite e no PostgreSQL): sem `lr` ou nulo
    e sem filtro; "lang_pt" exato e so portugues; qualquer outro valor, JSON invalido ou que nao e objeto vale "outro";
    chave repetida vale a ultima, como no json.loads."""
    import sqlite3
    con = sqlite3.connect(banco)
    con.execute("UPDATE chamadas SET parametros_json = ? WHERE chamada_id = (SELECT o.chamada_id FROM ocorrencias o "
                "JOIN artigos a ON a.artigo_id = o.artigo_id WHERE a.link = 'https://s.org/LA4')", (parametros,))
    con.commit()
    con.close()
    assert _idioma_de(banco, "https://s.org/LA4") == chave


def test_chaves_do_artigo_em_ordem_alfabetica_sem_repeticao(banco):
    assert _idioma_de(banco, "https://s.org/LA2") == "sem_filtro,so_portugues"
    assert _idioma_de(banco, "https://s.org/LA1") == "so_portugues"
    assert "chamadas" not in LeituraPainelSQLite(banco).linhas_artigos_das_consultas(list(CONJUNTO))[0]


def test_filtro_de_idioma_com_a_fonte_web_fica_vazio(banco, tmp_path):
    html = _cli(banco, tmp_path).get("/literatura/referencias?fonte=Web&idioma=so_portugues").text
    assert "<strong>0</strong> referência(s)" in html


def test_parametro_repetido_vale_o_ultimo(banco, tmp_path):
    html = _cli(banco, tmp_path).get("/literatura/referencias?idioma=as_duas&idioma=sem_filtro").text
    assert sorted(_idioma_por_url(html)) == ["https://s.org/LA3"] and '<option value="sem_filtro" selected>' in html


def test_rotulo_conta_as_buscas_de_qualquer_conjunto(banco, tmp_path):
    """Com o filtro de conjunto, o rotulo continua contando todas as buscas que acharam o artigo (como diz a legenda)."""
    html = _cli(banco, tmp_path).get("/literatura/referencias?conjunto=A&idioma=as_duas").text
    assert sorted(_idioma_por_url(html)) == ["https://s.org/LA2"]


def test_uma_consulta_a_mais_por_lote_de_chamadas(banco):
    """As chaves de idioma saem de uma consulta so as chamadas (em lotes de 500), nunca uma por artigo."""
    leitura = LeituraPainelSQLite(banco)
    feitas = []
    original = leitura._consulta
    leitura._consulta = lambda sql, args=(): feitas.append(sql) or original(sql, args)
    leitura.linhas_artigos_das_consultas(list(CONJUNTO))
    assert len(feitas) == 2 and "FROM chamadas WHERE chamada_id IN" in feitas[1]
