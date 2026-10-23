"""Tela Referências do banco `vies`: recorte e local vindos da triagem, consultas por extenso, escolha das colunas, os dois
downloads (modelo Érico e colunas selecionadas) com os filtros da tela, e as execuções clicáveis da visão geral."""
import io
import re

import openpyxl
import pytest
from fastapi.testclient import TestClient

from adaptadores.bruto_local import BrutoLocal
from adaptadores.leitura_sqlite import LeituraPainelSQLite
from adaptadores.repositorio_sqlite import RepositorioSQLite
from config import consultas as C, triagem_escopo as TE, vocabulario as V
from config.eixos_tematicos import EIXOS_EXPLORACAO
from dominio import referencias_busca as RB
from painel.app import Contexto, criar_app
from test_triagem_vies import construir_vies
from test_triagem_vies_painel import _linha, _triar

CONJUNTO = {t.id: n for n, ts in C.CONJUNTOS_LITERATURA.items() for t in ts}
TEXTOS = {t.id: t.query for ts in C.CONJUNTOS_LITERATURA.values() for t in ts}
ROTULOS = dict(TE.RECORTES_VIES)


# ---------- colunas e linhas ----------
def test_colunas_da_tela_na_ordem():
    assert RB.COLUNAS_TELA == RB.COLUNAS + ("Recorte", "Local", "Consultas", "Execução", "Fonte da busca", "Conjunto")
    assert RB.LOCAIS == {"brasil": "Brasil", "outro": "Internacional", "indeterminado": "Indeterminado"}


def _linha_tela():
    return RB.linha_web({"referencia_id": 1, "dominio": "e.org", "data_pub": "2024-01-01", "titulo": "T", "url_original": "https://e.org",
                         "conjuntos": "C", "execucoes": "r1", "consultas": "LC2,LC1"})


@pytest.mark.parametrize("triagem, recorte, local", [
    ({"valido": 1, "recorte": "vies_algoritmico", "pais_escopo": "Brasil"}, "Viés algorítmico", "Brasil"),
    ({"valido": 1, "recorte": "aplicacoes_tecnologicas", "pais_escopo": "EUA / Brasil"}, "Aplicações tecnológicas", "Brasil"),
    ({"valido": 1, "recorte": "aplicacoes_tecnologicas", "pais_escopo": "Reino Unido"}, "Aplicações tecnológicas", "Internacional"),
    ({"valido": 1, "recorte": None, "pais_escopo": "Indeterminado"}, "", "Indeterminado"),        # linha da v1: sem recorte
    ({"valido": 1, "recorte": "outro", "pais_escopo": None}, "", "Indeterminado"),
    ({"valido": 1, "pais_escopo": "Brasil"}, "", "Brasil"),                                            # coluna ausente
    ({"valido": 1, "recorte": "vies_algoritmico", "pais_escopo": "  "}, "Viés algorítmico", "Internacional"),  # regra de categoria_pais
    ({"valido": 0, "recorte": "vies_algoritmico", "pais_escopo": "Brasil"}, "", ""),                   # triagem inválida
    (None, "", ""),
])
def test_recorte_e_local_vem_da_triagem_valida(triagem, recorte, local):
    linha = RB.com_triagem(_linha_tela(), triagem, ROTULOS)
    assert (linha["Recorte"], linha["Local"]) == (recorte, local)


def test_consultas_da_linha_sao_os_codigos_em_ordem():
    assert _linha_tela()["Consultas"] == ["LC1", "LC2"]


def test_filtro_do_recorte_e_o_rotulo_exato():
    linhas = [{**_linha_tela(), "Recorte": r, "ID": f"W-{i}"} for i, r in enumerate(["Viés algorítmico", "Aplicações tecnológicas", ""], 1)]
    assert [l["ID"] for l in RB.filtrar(linhas, recorte="Aplicações tecnológicas")] == ["W-2"]
    assert [l["ID"] for l in RB.filtrar(linhas)] == ["W-1", "W-2", "W-3"]


@pytest.mark.parametrize("pedidas, esperado", [
    ([], list(RB.COLUNAS_TELA)),
    (["Título", "ID"], ["ID", "Título"]),                        # a ordem é sempre a da tela
    (["ID", "ID", "Recorte"], ["ID", "Recorte"]),                 # repetição ignorada
    (["ID", "xyz"], None), ([""], None), (["id"], None), ([" ID"], None),
    (["T\u0131tulo"], None), (["Ti\u0301tulo"], None),                        # outra letra e a forma NFD: so o nome exato
])
def test_colunas_escolhidas(pedidas, esperado):
    assert RB.colunas_escolhidas(pedidas) == esperado


@pytest.mark.parametrize("coluna, valor, esperado", [
    ("Consultas", ["LC1", "ZZ9"], "inteligência artificial segurança pública; ZZ9"),
    ("Consultas", [], ""), ("Execução", ["r1", "r2"], "r1, r2"), ("Execução", [], ""),
    ("Ano", 2024, "2024"), ("Ano", None, ""), ("Título", "T", "T"),
])
def test_valor_texto(coluna, valor, esperado):
    assert RB.valor_texto({coluna: valor}, coluna, TEXTOS) == esperado


def test_textos_das_consultas_e_larguras():
    assert RB.textos_das_consultas(["LC1", "ZZ9"], TEXTOS) == [TEXTOS["LC1"], "ZZ9"] and RB.textos_das_consultas([], TEXTOS) == []
    assert set(RB.LARGURA_DA_COLUNA) == set(RB.COLUNAS_TELA)


# ---------- tela e downloads ----------
@pytest.fixture
def banco(tmp_path):
    construir_vies(tmp_path / "busca.db", tmp_path / "raw")
    import sqlite3
    con = sqlite3.connect(tmp_path / "busca.db")
    refs = [r[0] for r in con.execute("SELECT referencia_id FROM referencias_web ORDER BY 1")]
    con.close()
    _triar(tmp_path / "busca.db", [_linha("referencia", refs[0], recorte="aplicacoes_tecnologicas", pais_escopo="EUA")]
           + [_linha("referencia", i, recorte="vies_algoritmico") for i in refs[1:]])
    return tmp_path / "busca.db"


def _cli(banco, tmp_path):
    ctx = Contexto(banco="vies", nome_banco="busca.db", somente_leitura=True, pasta_backup=tmp_path / "b", ufs=dict(C.UFS),
                   status_permitidos=set(V.STATUS), status_definicoes=dict(V.STATUS), hosts_permitidos=("testserver",),
                   conjunto_da_consulta=dict(CONJUNTO), consultas_literatura=dict(TEXTOS), eixos_exploracao=dict(EIXOS_EXPLORACAO),
                   versao_triagem_vies=TE.PROMPT_TRIAGEM_VIES_VERSAO, versoes_triagem_vies=TE.VERSOES_TRIAGEM_VIES,
                   recortes_vies=dict(ROTULOS))
    return TestClient(criar_app(LeituraPainelSQLite(banco), lambda: RepositorioSQLite(banco), BrutoLocal(tmp_path), ctx))


def _cabecalhos(html):
    return re.findall(r"<th(?: class=\"num\")?>([^<]+)</th>", html)


def test_tela_mostra_recorte_local_e_consultas_por_extenso(banco, tmp_path):
    html = _cli(banco, tmp_path).get("/literatura/referencias").text
    assert _cabecalhos(html) == list(RB.COLUNAS_TELA)
    assert "<td>Aplicações tecnológicas</td><td>Internacional</td>" in html and "<td>Viés algorítmico</td><td>Brasil</td>" in html
    assert f"<li>{TEXTOS['LA1']}</li>" in html and "<td>LA1" not in html and '<ul class="consultas-lista">' in html
    assert '<option value="aplicacoes_tecnologicas">Aplicações tecnológicas (sem foco no viés)</option>' in html


def test_filtro_do_recorte_na_tela_e_nos_dois_downloads(banco, tmp_path):
    cli = _cli(banco, tmp_path)
    assert "<strong>1</strong> referência(s)" in cli.get("/literatura/referencias?recorte=aplicacoes_tecnologicas").text
    assert "<strong>3</strong> referência(s)" in cli.get("/literatura/referencias?recorte=vies_algoritmico").text
    erico = openpyxl.load_workbook(io.BytesIO(cli.get("/literatura/referencias.xlsx?recorte=aplicacoes_tecnologicas&modelo=erico").content))
    assert erico.sheetnames == ["Referências ampliadas", "Esquema de exploração"] and erico.active.max_row == 2
    colunas = openpyxl.load_workbook(io.BytesIO(cli.get("/literatura/referencias.xlsx?recorte=vies_algoritmico&modelo=colunas").content))
    assert colunas.sheetnames == ["Referências"] and colunas.active.max_row == 4


@pytest.mark.parametrize("consulta, status", [
    ("recorte=vies", 404), ("recorte=Viés+algorítmico", 404), ("recorte=%20vies_algoritmico", 404), ("colunas=xyz", 404),
    ("colunas=", 404), ("colunas=ID&colunas=nada", 404), ("colunas=Ti%CC%81tulo", 404),
    ("modelo=erico", 200), ("modelo=zz", 200),                              # a tela ignora o modelo
    ("recorte=xx&recorte=vies_algoritmico", 200), ("recorte=vies_algoritmico&recorte=xx", 404),   # repetido: vale o ultimo
])
def test_filtro_ou_coluna_invalida_da_404_na_tela(banco, tmp_path, consulta, status):
    assert _cli(banco, tmp_path).get(f"/literatura/referencias?{consulta}").status_code == status


@pytest.mark.parametrize("consulta, status", [
    ("modelo=erico", 200), ("modelo=colunas", 200), ("modelo=", 404), ("modelo=Erico", 404), ("modelo=outro", 404),
    ("modelo=xx&modelo=erico", 200), ("modelo=erico&modelo=xx", 404),    # repetido: vale o ultimo
    ("colunas=xyz&modelo=erico", 404),                                    # o modelo Erico ignora as colunas, mas recusa a invalida
])
def test_modelo_do_download(banco, tmp_path, consulta, status):
    assert _cli(banco, tmp_path).get(f"/literatura/referencias.xlsx?{consulta}").status_code == status


def test_modelo_erico_ignora_as_colunas_escolhidas(banco, tmp_path):
    wb = openpyxl.load_workbook(io.BytesIO(_cli(banco, tmp_path).get("/literatura/referencias.xlsx?colunas=ID&modelo=erico").content))
    assert [c.value for c in wb.active[1]] == list(RB.COLUNAS)


def test_download_de_todas_as_colunas_formatado_como_o_modelo(banco, tmp_path):
    cli = _cli(banco, tmp_path)
    wb = openpyxl.load_workbook(io.BytesIO(cli.get("/literatura/referencias.xlsx?modelo=colunas").content))
    erico = openpyxl.load_workbook(io.BytesIO(cli.get("/literatura/referencias.xlsx?modelo=erico").content)).active
    f = wb.active
    assert [c.value for c in f[1]] == list(RB.COLUNAS_TELA)
    assert [f.column_dimensions[openpyxl.utils.get_column_letter(i)].width for i in range(1, 18)] == [RB.LARGURA_DA_COLUNA[c] for c in RB.COLUNAS_TELA]
    assert f.freeze_panes == "A2" and f.auto_filter.ref == f"A1:Q{f.max_row}"
    assert f["A1"].fill.fgColor.rgb == erico["A1"].fill.fgColor.rgb and f["A1"].font.b == erico["A1"].font.b
    assert f["I2"].value is None                                          # texto vazio vira celula vazia


def test_download_sem_linhas(banco, tmp_path):
    wb = openpyxl.load_workbook(io.BytesIO(_cli(banco, tmp_path).get("/literatura/referencias.xlsx?q=zzzz&modelo=colunas&colunas=ID").content))
    assert wb.active.max_row == 1 and [c.value for c in wb.active[1]] == ["ID"]


def test_sem_modelo_baixa_o_modelo_erico(banco, tmp_path):
    wb = openpyxl.load_workbook(io.BytesIO(_cli(banco, tmp_path).get("/literatura/referencias.xlsx").content))
    assert wb.sheetnames == ["Referências ampliadas", "Esquema de exploração"]
    assert [c.value for c in wb.active[1]] == list(RB.COLUNAS)


def test_colunas_escolhidas_na_tela_na_paginacao_e_no_download(banco, tmp_path):
    cli = _cli(banco, tmp_path)
    html = cli.get("/literatura/referencias?colunas=Título&colunas=ID&colunas=Consultas&colunas=Local").text
    assert _cabecalhos(html) == ["ID", "Título", "Local", "Consultas"]
    assert "Colunas (4 de 17)" in html and 'value="Recorte">' in html and 'value="ID" checked>' in html
    base = "colunas=ID&amp;colunas=T%C3%ADtulo&amp;colunas=Local&amp;colunas=Consultas"
    assert f'href="/literatura/referencias.xlsx?{base}&amp;modelo=colunas">Colunas selecionadas</a>' in html
    assert f'href="/literatura/referencias.xlsx?{base}&amp;modelo=erico">Modelo Érico</a>' in html
    assert '<a href="/literatura/referencias">Limpar</a>' in html
    wb = openpyxl.load_workbook(io.BytesIO(cli.get(f"/literatura/referencias.xlsx?{base.replace('&amp;', '&')}&modelo=colunas").content))
    folha = wb.active
    assert [c.value for c in folha[1]] == ["ID", "Título", "Local", "Consultas"]
    linhas = [[c.value for c in r] for r in folha.iter_rows(min_row=2)]
    assert all(isinstance(v, str) for l in linhas for v in l if v is not None)
    assert any(TEXTOS["LA1"] in (l[3] or "") for l in linhas)


def test_todas_as_colunas_nao_vao_no_endereco(banco, tmp_path):
    todas = "&".join(f"colunas={c}" for c in RB.COLUNAS_TELA)
    explicitas = _cli(banco, tmp_path).get(f"/literatura/referencias?{todas}").text      # o formulario com tudo marcado
    assert '<a href="/literatura/referencias">Limpar</a>' not in explicitas and 'modelo=erico">' in explicitas
    assert "colunas=" not in explicitas.split("<table")[0].split("escolha-colunas-itens")[0]
    html = _cli(banco, tmp_path).get("/literatura/referencias").text
    assert 'href="/literatura/referencias.xlsx?modelo=erico">Modelo Érico</a>' in html and "colunas=" not in html.split("<table")[0].split("escolha-colunas-itens")[0]
    assert "Colunas (17 de 17)" in html and '<a href="/literatura/referencias">Limpar</a>' not in html


def test_menu_baixar_tem_os_dois_modelos(banco, tmp_path):
    html = _cli(banco, tmp_path).get("/literatura/referencias?conjunto=B").text
    form = re.search(r'<form class="filtros".*?</form>', html, re.S).group(0)
    assert '<details class="menu-baixar"><summary class="botao secundario"' in form and ">Baixar</summary>" in form
    assert form.index(">Modelo Érico</a>") < form.index(">Colunas selecionadas</a>")
    assert 'href="/literatura/referencias.xlsx?conjunto=B&amp;modelo=erico"' in form


# ---------- visão geral: execuções clicáveis ----------
def test_ultimas_execucoes_levam_a_execucao(banco, tmp_path):
    html = _cli(banco, tmp_path).get("/").text
    run = LeituraPainelSQLite(banco).execucoes_recentes(1)[0]["run_id"]
    assert f'<td><a href="/execucoes/{run}"><code>{run}</code></a></td>' in html


def test_ultimas_execucoes_sem_permissao_sao_texto(banco, tmp_path):
    from adaptadores.contas_sqlite import ContasSQLite
    from dominio import autenticacao as A
    contas = ContasSQLite(tmp_path / "contas.db")
    contas.garantir_padrao()
    grupos = {g["nome"]: g["grupo_id"] for g in contas.grupos()}
    contas.criar_usuario("leitor", "leitor", grupos["Leitor"], A.hash_senha("senha-de-teste-123"), "2026-10-22T00:00:00+00:00")
    ctx = Contexto(banco="vies", nome_banco="busca.db", somente_leitura=False, pasta_backup=tmp_path / "b", ufs=dict(C.UFS),
                   status_permitidos=set(V.STATUS), status_definicoes=dict(V.STATUS), hosts_permitidos=("testserver",),
                   conjunto_da_consulta=dict(CONJUNTO), eixos_exploracao=dict(EIXOS_EXPLORACAO), autenticacao=True,
                   versao_triagem_vies=TE.PROMPT_TRIAGEM_VIES_VERSAO, recortes_vies=dict(ROTULOS))
    cli = TestClient(criar_app(LeituraPainelSQLite(banco), lambda: RepositorioSQLite(banco), BrutoLocal(tmp_path), ctx, contas=contas))
    csrf = re.search(r'name="csrf" value="([^"]+)"', cli.get("/entrar").text).group(1)
    assert cli.post("/entrar", data={"login": "leitor", "senha": "senha-de-teste-123", "proximo": "/", "csrf": csrf},
                    follow_redirects=False).status_code == 303
    run = LeituraPainelSQLite(banco).execucoes_recentes(1)[0]["run_id"]
    html = cli.get("/").text
    assert f"<td><code>{run}</code></td>" in html and f'href="/execucoes/{run}"' not in html


def test_paginacao_leva_as_colunas(tmp_path):
    from aplicacao.comum import registrar_chamada
    from dominio.modelos import RegistroChamada, ResultadoWeb
    from infra.util import agora
    caminho = tmp_path / "muitas.db"
    repo = RepositorioSQLite(caminho)
    try:
        run = repo.iniciar_execucao("buscar-literatura", "h", "{}", "c")
        reg = RegistroChamada(provedor="f", operacao="buscar", parametros={}, requisitado_em=agora(), http_status=200)
        cid, _ = registrar_chamada(repo, None, run, "busca_literatura", "consulta_literatura", "LA1", "k", reg, [])
        repo.salvar_referencias_web(cid, run, "A", "LA1", [ResultadoWeb(f"https://e.org/{i}", f"t{i}", "", None, i) for i in range(1, 61)])
        repo.confirmar()
    finally:
        repo.fechar()
    html = _cli(caminho, tmp_path).get("/literatura/referencias?colunas=ID").text
    assert 'href="/literatura/referencias?colunas=ID&amp;p=2">Próxima' in html


def test_visao_geral_do_banco_real_tambem_tem_as_execucoes_clicaveis(banco, tmp_path):
    ctx = Contexto(banco="real", nome_banco="busca.db", somente_leitura=True, pasta_backup=tmp_path / "b", ufs=dict(C.UFS),
                   status_permitidos=set(V.STATUS), status_definicoes=dict(V.STATUS), hosts_permitidos=("testserver",))
    html = TestClient(criar_app(LeituraPainelSQLite(banco), lambda: RepositorioSQLite(banco), BrutoLocal(tmp_path), ctx)).get("/").text
    run = LeituraPainelSQLite(banco).execucoes_recentes(1)[0]["run_id"]
    assert f'<td><a href="/execucoes/{run}"><code>{run}</code></a></td>' in html
