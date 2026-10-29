"""Observacoes, motivos das decisoes humanas e motivos das sugestoes so saem do servidor para quem tem `ver_registros_revisao`
ou pode decidir aquele item; o visitante anonimo nunca os recebe. As planilhas exigem todas as permissoes dos dados que levam.
A tabela agregada de motivos de exclusao do PRISMA continua publica (decisao do pesquisador)."""
import io
import re
import shutil
import sqlite3

import pytest
from apoio_divulgacao import gravar_politica
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from adaptadores.bruto_local import BrutoLocal
from adaptadores.contas_sqlite import ContasSQLite
from adaptadores.leitura_sqlite import LeituraPainelSQLite
from adaptadores.repositorio_sqlite import RepositorioSQLite
from config import consultas as C, vocabulario as V
from config.modelos_exportacao import MODELOS_REFERENCIAS
from dominio import autenticacao as A
from painel.app import Contexto, criar_app
from test_autenticacao import _criar, _csrf, _entrar, painel  # noqa: F401
from test_links_por_permissao import SEMENTES_REAL, _corpo, _percorrer
from test_painel import demo  # noqa: F401
from test_contas_postgres import pg  # noqa: F401

OBS, REV, SUG, MOT = "SENTINELA-OBS-7Q", "SENTINELA-REV-7Q", "SENTINELA-SUG-7Q", "SENTINELA-MOT-7Q"
NOVAS = {"ver_registros_revisao", "ver_custos"}
EXPORTACOES = {
    "/execucoes/resultado.xlsx": ("ver_chamadas", "ver_nomes_revisores", "ver_custos", "ver_registros_revisao"),
    "/literatura/catalogo.xlsx": ("triar_artigos", "ver_registros_revisao"),
    "/achados/casos.xlsx": ("verificar_achados", "ver_registros_revisao"),
}
# grupos criados no teste: um "Leitor" e um "Revisor" antigos, ainda sem as permissoes novas, e quem decide sem ver registros
GRUPOS_EXTRAS = {
    "Consulta": {"ver_painel", "ver_nomes_revisores"},
    "Tecnico": {"ver_painel", "ver_nomes_revisores", "ver_chamadas", "verificar_achados", "triar_artigos", "triar_texto_completo"},
    "Decisor": {"ver_painel", "verificar_achados", "triar_artigos"},
}


# ---------------- catalogo ----------------
def test_permissoes_novas_no_catalogo():
    assert A.PERMISSOES["ver_registros_revisao"] == "Ver observações e motivos das decisões"
    assert A.PERMISSOES["ver_custos"] == "Ver custos das execuções"
    assert list(A.PERMISSOES)[-2:] == ["ver_registros_revisao", "ver_custos"]
    assert A.CATALOGO_PERMISSOES["ver_registros_revisao"] == {"modulo": "consulta", "sensibilidade": "dados_revisao",
                                                              "depende_de": ("ver_painel",)}
    assert A.CATALOGO_PERMISSOES["ver_custos"] == {"modulo": "tecnico", "sensibilidade": "financeira", "depende_de": ("ver_painel",)}
    assert A.SENSIBILIDADES[-2:] == ("dados_revisao", "financeira")
    assert A.ROTULOS_SENSIBILIDADE["dados_revisao"] == "dados de revisão" and A.ROTULOS_SENSIBILIDADE["financeira"] == "financeira"
    assert NOVAS <= A.PERMISSOES_PROIBIDAS_VISITANTE and "ver_chamadas" not in A.PERMISSOES_PROIBIDAS_VISITANTE
    assert NOVAS <= A.GRUPOS_PADRAO["Leitor"] and NOVAS <= A.GRUPOS_PADRAO["Revisor"]
    assert not NOVAS & A.GRUPOS_PADRAO["Visitante"]
    assert A.EXIGENCIAS_EXPORTACOES == EXPORTACOES


@pytest.mark.parametrize("caminho", sorted(EXPORTACOES))
@pytest.mark.parametrize("metodo", ["GET", "HEAD"])
def test_exportacao_exige_todas_as_permissoes(caminho, metodo):
    exige = set(EXPORTACOES[caminho])
    assert A.pode_acessar("Qualquer", exige, metodo, caminho) is True
    for falta in sorted(exige):
        assert A.pode_acessar("Qualquer", exige - {falta}, metodo, caminho) is False, falta
    assert A.pode_acessar(A.GRUPO_ADMIN, set(), metodo, caminho) is True
    assert A.pode_acessar("Qualquer", exige, "POST", caminho) is False          # POST nessas rotas continua do administrador


def test_outras_rotas_nao_mudam():
    assert A.pode_acessar("Leitor", {"ver_painel"}, "GET", "/literatura/referencias.xlsx") is True
    assert A.pode_acessar("Leitor", {"ver_painel"}, "GET", "/achados/7") is True
    assert A.pode_acessar("Leitor", {"ver_painel"}, "GET", "/execucoes") is False


# ---------------- adaptador: o visitante nunca recebe as novas ----------------
@pytest.mark.parametrize("nova", sorted(NOVAS))
def test_visitante_nao_recebe_as_novas(painel, nova):  # noqa: F811
    _, contas, _, _ = painel
    vis = next(g for g in contas.grupos() if g["nome"] == A.GRUPO_VISITANTE)
    antes = set(vis["permissoes"])
    with pytest.raises(ValueError):
        contas.definir_permissoes(vis["grupo_id"], {"ver_painel", nova})
    assert next(g for g in contas.grupos() if g["nome"] == A.GRUPO_VISITANTE)["permissoes"] == antes


@pytest.mark.postgres
@pytest.mark.parametrize("nova", sorted(NOVAS))
def test_visitante_nao_recebe_as_novas_no_postgres(pg, nova):  # noqa: F811
    from test_contas_postgres import _gid as _gid_pg
    c, _, _ = pg
    antes = c.permissoes_do_grupo(A.GRUPO_VISITANTE)
    with pytest.raises(ValueError, match=f"O grupo Visitante não pode receber essa permissão: {nova}"):
        c.definir_permissoes(_gid_pg(c, A.GRUPO_VISITANTE), {"ver_painel", nova})
    assert c.permissoes_do_grupo(A.GRUPO_VISITANTE) == antes
    assert NOVAS <= c.permissoes_do_grupo("Leitor") and NOVAS <= c.permissoes_do_grupo("Revisor")    # semeadura de banco novo


def test_semeadura_so_vale_para_grupos_novos(painel):  # noqa: F811
    _, contas, _, contas_db = painel
    leitor = next(g for g in contas.grupos() if g["nome"] == "Leitor")
    assert NOVAS <= leitor["permissoes"]                          # banco de contas novo: semeado com as novas
    contas.definir_permissoes(leitor["grupo_id"], {"ver_painel"})
    contas.garantir_padrao()                                      # ressemear nao devolve nada a um grupo existente
    assert next(g for g in contas.grupos() if g["nome"] == "Leitor")["permissoes"] == {"ver_painel"}


# ---------------- painel: cenario com textos sentinela gravados pelas rotas ----------------
def _form(cli, url):
    h = cli.get(url).text
    return h, _csrf(cli)


def _campo(h, nome):
    m = re.search(rf'name="{nome}" value="([^"]*)"', h)
    return m.group(1) if m else ""


@pytest.fixture
def painel_modelos(demo, tmp_path):  # noqa: F811
    """O painel do banco demo com autenticacao e com os modelos de planilha do catalogo (como o `painel` de test_autenticacao)."""
    banco, raw = demo
    copia = tmp_path / "demo.db"
    shutil.copy(banco, copia)
    ctx = Contexto(banco="real", nome_banco=copia.name, somente_leitura=False, pasta_backup=tmp_path / "backups",
                   ufs=dict(C.UFS), status_permitidos=set(V.STATUS) | set(V.STATUS_FINAL_EXTRA),
                   status_definicoes=dict(V.STATUS), hosts_permitidos=("testserver",), autenticacao=True,
                   modelos_referencias=dict(MODELOS_REFERENCIAS))
    contas_db = tmp_path / "contas.db"
    contas = ContasSQLite(contas_db)
    gravar_politica(contas, "real")
    app = criar_app(LeituraPainelSQLite(copia), lambda: RepositorioSQLite(copia), BrutoLocal(raw), ctx, contas=contas)
    return TestClient(app), contas, copia, contas_db


@pytest.fixture
def cenario(painel_modelos):
    painel = painel_modelos
    cli0, contas, copia, contas_db = painel
    for nome, perms in GRUPOS_EXTRAS.items():
        gid = contas.criar_grupo(nome, "2026-10-27T00:00:00+00:00")
        contas.definir_permissoes(gid, perms)
    logins = {"Administrador": "adm", "Revisor": "maria", "Leitor": "leo", "Consulta": "cons", "Tecnico": "tec", "Decisor": "dec"}
    clis = {"Visitante": cli0}
    for grupo, login in logins.items():
        _criar(painel, login, grupo, nome=login.title())
        c = TestClient(cli0.app)
        assert _entrar(c, login).status_code == 303
        clis[grupo] = c
    con = sqlite3.connect(copia)
    a1, a2 = [r[0] for r in con.execute("SELECT achado_id FROM v_achados_revisao WHERE verificado=0 ORDER BY achado_id LIMIT 2")]
    con.close()
    maria, adm = clis["Revisor"], clis["Administrador"]
    art = int(re.search(r'href="/literatura/artigos/(\d+)"', adm.get("/literatura/artigos").text).group(1))
    for a in (a1, a2):                                                       # maria verifica com observacao
        h, csrf = _form(maria, f"/achados/{a}")
        r = maria.post(f"/achados/{a}/verificar", data={"status_final": "em_uso", "observacao": OBS, "versao": _campo(h, "versao"),
                                                        "csrf": csrf}, follow_redirects=False)
        assert r.status_code == 303, r.text[:300]
    h, csrf = _form(adm, f"/achados/{a2}")                                  # o administrador revoga a2 com motivo
    form = h[h.index(f'action="/achados/{a2}/revogar'):]
    r = adm.post(f"/achados/{a2}/revogar", data={"motivo": REV, "versao": _campo(form, "versao"), "csrf": csrf, "revisor": "Adm"},
                 follow_redirects=False)
    assert r.status_code == 303, r.text[:300]
    h, csrf = _form(adm, f"/achados/{a1}")                                  # e sugere outro status em a1, com motivo
    estado = re.search(r'<input type="hidden" name="estado" value="([^"]+)">', h).group(1)
    r = adm.post(f"/achados/{a1}/sugerir", data={"valor": "piloto", "motivo": SUG, "estado": estado, "csrf": csrf}, follow_redirects=False)
    assert r.status_code == 303, r.text[:300]
    h, csrf = _form(maria, f"/literatura/artigos/{art}")                    # maria exclui um artigo com motivo
    r = maria.post(f"/literatura/artigos/{art}/triar", data={"decisao": "excluir", "motivo": MOT, "versao": _campo(h, "versao"),
                                                             "etapa": "titulo_resumo", "csrf": csrf}, follow_redirects=False)
    assert r.status_code == 303, r.text[:300]
    return {"clis": clis, "a1": a1, "a2": a2, "art": art}


VE_REGISTROS = {"Administrador": True, "Revisor": True, "Leitor": True, "Consulta": False, "Tecnico": True, "Decisor": True,
                "Visitante": False}            # Tecnico e Decisor: podem decidir o item, entao veem os registros do item


def test_paginas_do_item_so_com_a_permissao_ou_quem_decide(cenario):
    c = cenario
    for grupo, cli in c["clis"].items():
        if grupo == "Visitante":                                 # fechadas para quem nao entrou (D145)
            for u in (f"/achados/{c['a1']}", f"/achados/{c['a2']}", f"/literatura/artigos/{c['art']}"):
                assert cli.get(u, follow_redirects=False).status_code == 303, u
            continue
        a1, a2, art = (_corpo(cli.get(u).text) for u in (f"/achados/{c['a1']}", f"/achados/{c['a2']}", f"/literatura/artigos/{c['art']}"))
        ve = VE_REGISTROS[grupo]
        assert (OBS in a1) is ve and (SUG in a1) is ve, grupo
        assert (OBS in a2) is ve and (REV in a2) is ve, grupo
        assert (MOT in art) is ve, grupo


def test_quem_so_consulta_nao_recebe_registro_em_pagina_nenhuma(cenario):
    c = cenario
    extras = [f"/achados/{c['a1']}", f"/achados/{c['a2']}", f"/literatura/artigos/{c['art']}", f"/literatura/artigos/{c['art']}?ok=1"]
    for grupo in ("Visitante", "Consulta"):
        paginas = _percorrer(c["clis"][grupo], extras + SEMENTES_REAL, limite=120)
        if grupo == "Visitante":                           # quem nao entrou so abre as paginas publicas (D145)
            assert set(paginas) <= {"/", "/sobre"} and paginas
        else:
            assert len(paginas) >= 15
        cli = c["clis"][grupo]
        if grupo == "Visitante":                                 # so as paginas publicas abrem; o resto vai para a entrada
            for url in extras + ["/achados", "/literatura/artigos", "/literatura", "/buscar-texto?q=reconhecimento"]:
                assert cli.get(url, follow_redirects=False).status_code == 303, url
        vazados = []
        for url in extras + ["/", "/achados", "/literatura/artigos", "/literatura", "/buscar-texto?q=reconhecimento"]:
            h = cli.get(url).text
            vazados += [(url, s) for s in (OBS, REV, SUG, MOT) if s in h]
        assert vazados == [], grupo
        # o percurso inteiro, menos o PRISMA (a tabela agregada de motivos de exclusao e publica)
        for chave in paginas:
            if chave == "/literatura/prisma":
                continue
            h = cli.get(chave if "{" not in chave else extras[0]).text
            assert not any(s in h for s in (OBS, REV, SUG, MOT)), (grupo, chave)


def test_prisma_continua_mostrando_os_motivos_de_exclusao(cenario):
    for grupo in ("Consulta", "Administrador"):
        assert MOT in cenario["clis"][grupo].get("/literatura/prisma").text, grupo
    # quem nao entrou: o PRISMA fica fechado so ate a retirada dos numeros do modelo (fatia seguinte da D145)
    assert cenario["clis"]["Visitante"].get("/literatura/prisma", follow_redirects=False).status_code == 303


def test_formulario_de_quem_decide_traz_o_valor_atual(cenario):
    c = cenario
    h = c["clis"]["Decisor"].get(f"/achados/{c['a1']}").text
    assert re.search(rf'<textarea id="observacao"[^>]*>{OBS}</textarea>', h)
    h = c["clis"]["Decisor"].get(f"/literatura/artigos/{c['art']}").text
    assert re.search(rf'<textarea id="motivo"[^>]*>{MOT}</textarea>', h)


# ---------------- exportacoes ----------------
def _abre(cli, url):
    return cli.get(url, follow_redirects=False)


def _texto_xlsx(conteudo):
    wb = load_workbook(io.BytesIO(conteudo), read_only=True)
    return "\n".join(str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v is not None)


@pytest.mark.parametrize("caminho", sorted(EXPORTACOES))
def test_exportacoes_por_grupo(cenario, caminho):
    url = caminho + ("?modelo=padrao" if "catalogo" in caminho else "")
    for grupo, cli in cenario["clis"].items():
        perms = A.GRUPOS_PADRAO.get(grupo, GRUPOS_EXTRAS.get(grupo, set()))
        pode = grupo == A.GRUPO_ADMIN or set(EXPORTACOES[caminho]) <= set(perms)
        r = _abre(cli, url)
        assert (r.status_code == 200) is pode, (grupo, r.status_code)
        if not pode:
            assert r.status_code in (403, 303), (grupo, r.status_code)
        assert (cli.head(url, follow_redirects=False).status_code not in (403, 303)) is pode, grupo     # HEAD: autorizacao do GET
    assert _abre(cenario["clis"]["Tecnico"], url).status_code == 403          # o "Revisor" antigo, sem as novas, agora e recusado


def test_casos_xlsx_leva_a_observacao_para_quem_pode(cenario):
    r = cenario["clis"]["Revisor"].get("/achados/casos.xlsx")
    assert r.status_code == 200 and OBS in _texto_xlsx(r.content)


def test_rotas_das_planilhas_conferem_a_mesma_constante(cenario, monkeypatch):
    monkeypatch.setitem(A.EXIGENCIAS_EXPORTACOES, "/achados/casos.xlsx", ("verificar_achados", "ver_registros_revisao", "ver_custos"))
    sem_custos = cenario["clis"]["Decisor"]
    assert sem_custos.get("/achados/casos.xlsx", follow_redirects=False).status_code == 403
    assert cenario["clis"]["Revisor"].get("/achados/casos.xlsx", follow_redirects=False).status_code == 200


@pytest.mark.parametrize("caminho,pagina", [("/execucoes/resultado.xlsx", "/execucoes"), ("/achados/casos.xlsx", "/achados"),
                                            ("/literatura/catalogo.xlsx", "/literatura")])
def test_botoes_das_planilhas_so_para_quem_baixa(cenario, caminho, pagina):
    for grupo, cli in cenario["clis"].items():
        r = cli.get(pagina, follow_redirects=False)
        if r.status_code != 200:
            continue
        perms = A.GRUPOS_PADRAO.get(grupo, GRUPOS_EXTRAS.get(grupo, set()))
        pode = grupo == A.GRUPO_ADMIN or set(EXPORTACOES[caminho]) <= set(perms)
        assert (f'href="{caminho}' in _corpo(r.text)) is pode, (grupo, pagina)


def test_botao_de_cada_planilha_passa_pela_regra():
    from pathlib import Path
    achados = 0
    for arq in sorted(Path("painel/templates").glob("*.html")):
        s = arq.read_text(encoding="utf-8")
        for caminho in EXPORTACOES:
            for m in re.finditer(re.escape(f'href="{caminho}'), s):
                antes = s[:m.start()]
                se = antes.rfind("{% if")
                condicao = antes[se:antes.index("%}", se)]
                assert f'pode_seguir("{caminho}")' in condicao and "{% endif %}" not in antes[se:], (arq.name, caminho)
                achados += 1
    assert achados == 4          # resultado.xlsx, casos.xlsx e os dois modelos do catalogo


# ---------------- decisoes da pre-revisao: o que chega ao modelo, a mesma constante nas tres rotas, o texto da recusa ----------------
SEM_PERMISSAO = "Seu grupo não tem permissão para esta página ou ação."


@pytest.fixture
def capturas(monkeypatch):
    """O contexto que cada `pagina(...)` entrega ao modelo, e os objetos que a leitura devolveu (com uma copia do momento)."""
    import copy
    from fastapi.templating import Jinja2Templates
    contextos, devolvidos = [], []
    original = Jinja2Templates.TemplateResponse

    def espia(self, request, nome, contexto, *a, **k):
        contextos.append((nome, contexto))
        return original(self, request, nome, contexto, *a, **k)
    monkeypatch.setattr(Jinja2Templates, "TemplateResponse", espia)
    for metodo in ("achado", "historico_achado", "artigo", "sugestoes_abertas"):
        orig = getattr(LeituraPainelSQLite, metodo)

        def guarda(self, *a, _orig=orig, **k):
            r = _orig(self, *a, **k)
            devolvidos.append((r, copy.deepcopy(r)))
            return r
        monkeypatch.setattr(LeituraPainelSQLite, metodo, guarda)
    return contextos, devolvidos


def _contexto(contextos, nome):
    return [c for n, c in contextos if n == nome][-1]


def test_registros_saem_do_contexto_e_a_leitura_fica_intacta(cenario, capturas):
    contextos, devolvidos = capturas
    c, cons = cenario, cenario["clis"]["Consulta"]
    cons.get(f"/achados/{c['a2']}")
    ctx = _contexto(contextos, "achado.html")
    assert ctx["a"]["observacao"] is None and ctx["historico"] and all(h["observacao"] is None for h in ctx["historico"])
    cons.get(f"/achados/{c['a1']}")
    sugs = _contexto(contextos, "achado.html")["sugestoes"]
    assert sugs and all(s["motivo"] is None for s in sugs)
    cons.get(f"/literatura/artigos/{c['art']}")
    a = _contexto(contextos, "artigo.html")["a"]
    assert a["motivo"] is None and a["motivo_texto_completo"] is None and a["historico"]
    assert all(h["motivo"] is None and h["motivo_texto_completo"] is None for h in a["historico"])
    assert devolvidos and all(r == copia for r, copia in devolvidos)          # nada do que a leitura devolveu foi alterado
    assert any(OBS in repr(r) for r, _ in devolvidos) and any(MOT in repr(r) for r, _ in devolvidos)


def test_sugestao_muda_so_o_motivo(cenario, capturas):
    contextos, _ = capturas
    c = cenario
    c["clis"]["Revisor"].get(f"/achados/{c['a1']}")
    com = _contexto(contextos, "achado.html")["sugestoes"]
    c["clis"]["Consulta"].get(f"/achados/{c['a1']}")
    sem = _contexto(contextos, "achado.html")["sugestoes"]
    assert len(com) == len(sem) == 1 and com[0]["motivo"] == SUG and sem[0]["motivo"] is None
    fora = {"motivo", "minha"}                    # "minha" depende de quem ve (regra que ja existia), nao da restricao
    assert {k: v for k, v in com[0].items() if k not in fora} == {k: v for k, v in sem[0].items() if k not in fora}
    assert list(com[0]) == list(sem[0])


@pytest.mark.parametrize("caminho,reduzida,quem", [
    ("/execucoes/resultado.xlsx", ("ver_chamadas",), "Tecnico"),
    ("/achados/casos.xlsx", ("verificar_achados",), "Decisor"),
    ("/literatura/catalogo.xlsx", ("triar_artigos",), "Decisor"),
])
def test_cada_rota_le_a_constante(cenario, monkeypatch, caminho, reduzida, quem):
    url = caminho + ("?modelo=padrao" if "catalogo" in caminho else "")
    cli = cenario["clis"][quem]
    assert cli.get(url, follow_redirects=False).status_code == 403
    monkeypatch.setitem(A.EXIGENCIAS_EXPORTACOES, caminho, reduzida)   # o middleware e a rota leem a mesma entrada
    assert cli.get(url, follow_redirects=False).status_code == 200


@pytest.mark.parametrize("caminho", sorted(EXPORTACOES))
def test_recusa_das_planilhas_com_o_texto_de_sempre(cenario, caminho):
    url = caminho + ("?modelo=padrao" if "catalogo" in caminho else "")
    r = cenario["clis"]["Consulta"].get(url, follow_redirects=False)
    assert r.status_code == 403 and SEM_PERMISSAO in r.text


@pytest.mark.parametrize("caminho,quem_tem_tudo,falta,reduzida,quem_tem_a_reduzida", [
    ("/execucoes/resultado.xlsx", "Revisor", "enviar_coletas", ("ver_chamadas",), "Tecnico"),
    ("/achados/casos.xlsx", "Revisor", "enviar_coletas", ("verificar_achados",), "Decisor"),
    ("/literatura/catalogo.xlsx", "Revisor", "enviar_coletas", ("triar_artigos",), "Decisor"),
])
def test_a_propria_rota_confere_a_constante(cenario, monkeypatch, caminho, quem_tem_tudo, falta, reduzida, quem_tem_a_reduzida):
    """Com o middleware neutralizado, so a rota decide: uma exigencia a mais na constante recusa, uma a menos libera."""
    url = caminho + ("?modelo=padrao" if "catalogo" in caminho else "")
    monkeypatch.setattr(A, "pode_acessar", lambda *a, **k: True)
    monkeypatch.setitem(A.EXIGENCIAS_EXPORTACOES, caminho, EXPORTACOES[caminho] + (falta,))
    r = cenario["clis"][quem_tem_tudo].get(url, follow_redirects=False)
    assert r.status_code == 403 and SEM_PERMISSAO in r.text
    monkeypatch.setitem(A.EXIGENCIAS_EXPORTACOES, caminho, reduzida)
    assert cenario["clis"][quem_tem_a_reduzida].get(url, follow_redirects=False).status_code == 200


def test_banco_vazio_nao_oferece_a_planilha_completa(tmp_path):
    """Sem execucoes nao ha o que baixar: o botao continua escondido, como antes (a condicao `itens` permanece)."""
    from test_painel import _app
    RepositorioSQLite(tmp_path / "v.db").fechar()
    app, _ = _app(tmp_path / "v.db", tmp_path / "raw", tmp_path, somente_leitura=True)
    h = TestClient(app).get("/execucoes").text
    assert "Baixar a planilha completa" not in h
