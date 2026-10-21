"""Planilha completa de resultados pelo painel (RB30): as mesmas abas do comando `exportar`, lidas do banco do painel."""
import io
import re
import shutil
import sqlite3
import sys

import openpyxl
import pandas as pd
import pytest
from fastapi.testclient import TestClient

import main
from adaptadores.bruto_local import BrutoLocal
from adaptadores.consultas_exportacao import CONSULTAS_EXPORTACAO
from adaptadores.contas_sqlite import ContasSQLite
from adaptadores import leitura_sqlite
from adaptadores.leitura_sqlite import LeituraPainelSQLite
from adaptadores.repositorio_sqlite import RepositorioSQLite
from aplicacao import exportacao
from config import consultas as C, vocabulario as V
from dominio import autenticacao as A
from painel.app import Contexto, criar_app

SENHA = "senha-de-teste-123"
ABAS = ["Execucoes", "Funil", "Custos", "Falhas", "Descartes", "Dominios", "Busca_web", "Achados", "Resumo_UF_bruto", "Secretarias",
        "Municipios", "Chamadas", "Metricas_termo", "Artigos", "Artigos_ano", "Multi_termo", "Sobreposicao", "Triagem",
        "Triagem_escopo", "Buscas_manuais", "Manuais_resumo", "Termos_rendimento", "Termos_pares", "Termos_idioma",
        "Termos_pagina", "Curva_descoberta", "PRISMA", "Motivos_exclusao", "Registro_buscas", "Web_rendimento", "Web_modelos",
        "Web_posicao", "Web_consultas_vazias", "UF_tecnologia", "Precisao_LLM", "Correcoes_humanas", "Qualidade_classif",
        "Extracao", "Eficiencia", "Classif_invalidas", "Consistencia_evidencia", "Gasto_real", "Saldos", "Metricas_run",
        "Download_resultado", "Classif_resultado", "Achados_descartados", "Confianca_ajustada", "Auditoria_prefiltro", "Fontes",
        "Fontes_sobreposicao", "Termos_fonte", "Termos_total_estabilidade", "Ganho_contexto", "Revisores_kappa", "Discordancias",
        "Revisores_kappa_achados", "Discordancias_achados", "Auditoria_negativos", "Passadas", "Historico_triagem",
        "Historico_achados"]


@pytest.fixture(scope="module")
def demo(copiar_demo, tmp_path_factory):
    banco, raw = copiar_demo("resultado_painel")
    return banco, raw, tmp_path_factory.mktemp("resultado_painel")


def _contexto(tmp, autenticacao):
    return Contexto(banco="real", nome_banco="demo", somente_leitura=False, pasta_backup=tmp / "backups", ufs=dict(C.UFS),
                    status_permitidos=set(V.STATUS), status_definicoes=dict(V.STATUS), hosts_permitidos=("testserver",),
                    autenticacao=autenticacao)


def _cliente(demo, autenticacao=False, contas=None):
    banco, raw, tmp = demo
    return TestClient(criar_app(LeituraPainelSQLite(banco), lambda: RepositorioSQLite(banco), BrutoLocal(raw),
                                _contexto(tmp, autenticacao), contas=contas))


def _abas_xlsx(conteudo):
    wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True)
    return {n: [tuple(r) for r in wb[n].iter_rows(values_only=True)] for n in wb.sheetnames}


# ---------- as consultas ----------
def test_consultas_sao_as_mesmas_abas_de_antes_na_mesma_ordem():
    assert list(CONSULTAS_EXPORTACAO) == ABAS
    assert all(len(a) <= 31 for a in ABAS) and len(set(ABAS)) == len(ABAS)


def test_comando_e_painel_montam_as_mesmas_tabelas(demo):
    banco, _, _ = demo
    repo = RepositorioSQLite(banco)
    try:
        a = exportacao.montar_tabelas(repo, C.UFS)
    finally:
        repo.fechar()
    b = exportacao.montar_tabelas_painel(LeituraPainelSQLite(banco), C.UFS)
    assert list(a) == list(b) == [x for x in ABAS if x != "Resumo_UF_bruto"] + ["Resumo_UF"]
    for aba in a:
        pd.testing.assert_frame_equal(a[aba], b[aba], obj=aba)
    assert len(b["Achados"]) > 0 and len(b["Chamadas"]) > 0                  # a demonstracao tem dados


def test_leitura_do_painel_le_tudo_numa_so_transacao(demo, monkeypatch):
    banco, _, _ = demo
    comandos = []
    original = sqlite3.connect

    def espiao(*a, **k):
        con = original(*a, **k)
        con.set_trace_callback(comandos.append)
        return con
    conexoes = []
    original_espiao = espiao

    def contando(*a, **k):
        conexoes.append(a)
        return original_espiao(*a, **k)
    monkeypatch.setattr(sqlite3, "connect", contando)
    LeituraPainelSQLite(banco).tabelas_para_exportacao()
    assert len(conexoes) == 1                                               # uma conexao so
    assert comandos[:2] == ["PRAGMA busy_timeout = 5000", "BEGIN"]
    consultas = [n for n, c in enumerate(comandos) if c.lstrip().upper().startswith(("SELECT", "WITH"))]
    assert len(consultas) >= len(ABAS)
    meio = comandos[consultas[0]:consultas[-1] + 1]
    assert not [c for c in meio if c.strip().upper().startswith(("COMMIT", "END", "ROLLBACK", "BEGIN"))]


# ---------- a rota ----------
def test_download_sem_login_traz_as_mesmas_abas_e_valores_do_comando(demo, tmp_path, monkeypatch):
    banco, _, _ = demo
    r = _cliente(demo).get("/execucoes/resultado.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert re.fullmatch(r'attachment; filename="resultado_\d{8}_\d{4}\.xlsx"', r.headers["content-disposition"])
    assert r.headers["cache-control"] == "no-store"
    painel = _abas_xlsx(r.content)
    repo = RepositorioSQLite(banco)
    try:
        from adaptadores.exportador_excel import ExportadorExcel
        caminho = ExportadorExcel(tmp_path).exportar(exportacao.montar_tabelas(repo, C.UFS), tmp_path / "cli.xlsx")
    finally:
        repo.fechar()
    with open(caminho, "rb") as f:
        comando = _abas_xlsx(f.read())
    assert list(painel) == list(comando)
    assert painel == comando


@pytest.mark.parametrize("banco, prefixo", [("real", ""), ("simulado", "simulado_"), ("demo", "demo_")])
def test_nome_do_arquivo_marca_simulado_e_demo(demo, banco, prefixo):
    banco_arq, raw, tmp = demo
    ctx = _contexto(tmp, False)
    ctx.banco = banco
    cli = TestClient(criar_app(LeituraPainelSQLite(banco_arq), lambda: RepositorioSQLite(banco_arq), BrutoLocal(raw), ctx))
    r = cli.get("/execucoes/resultado.xlsx")
    assert re.fullmatch(rf'attachment; filename="{prefixo}resultado_\d{{8}}_\d{{4}}\.xlsx"', r.headers["content-disposition"])


def test_banco_ocupado_devolve_503_e_libera_a_vez(demo, monkeypatch):
    cli = _cliente(demo)

    def ocupado(conn):
        raise sqlite3.OperationalError("database is locked")
    monkeypatch.setattr(leitura_sqlite, "ler_tabelas", ocupado)            # o leitor traduz para BancoIndisponivel
    r = cli.get("/execucoes/resultado.xlsx")
    assert r.status_code == 503 and r.headers["content-type"].startswith("text/html")
    assert "O banco não respondeu a tempo para gerar a planilha completa. Tente de novo em alguns segundos." in r.text
    monkeypatch.undo()
    assert cli.get("/execucoes/resultado.xlsx").status_code == 200          # a vez foi liberada


def test_outra_planilha_em_preparo_devolve_503(demo, monkeypatch):
    import threading
    cli = _cliente(demo)
    dentro, solta = threading.Event(), threading.Event()
    original = exportacao.montar_tabelas_painel

    def devagar(*a, **k):
        dentro.set()
        solta.wait(30)
        return original(*a, **k)
    monkeypatch.setattr(exportacao, "montar_tabelas_painel", devagar)
    primeira = {}
    th = threading.Thread(target=lambda: primeira.setdefault("r", cli.get("/execucoes/resultado.xlsx")))
    th.start()
    assert dentro.wait(30)
    r = cli.get("/execucoes/resultado.xlsx")
    solta.set()
    th.join(60)
    assert r.status_code == 503 and "Outra planilha completa está sendo gerada agora. Tente de novo em alguns segundos." in r.text
    assert primeira["r"].status_code == 200


def test_leitor_traduz_o_erro_do_banco(demo, monkeypatch):
    from dominio.erros import BancoIndisponivel

    def ocupado(conn):
        raise sqlite3.OperationalError("database is locked")
    monkeypatch.setattr(leitura_sqlite, "ler_tabelas", ocupado)
    with pytest.raises(BancoIndisponivel):
        LeituraPainelSQLite(demo[0]).tabelas_para_exportacao()


def test_sem_banco_mostra_a_pagina_de_sem_banco(tmp_path):
    ctx = _contexto(tmp_path, False)
    cli = TestClient(criar_app(LeituraPainelSQLite(tmp_path / "nao_existe.db"), lambda: None, BrutoLocal(tmp_path), ctx))
    r = cli.get("/execucoes/resultado.xlsx")
    assert r.status_code == 503 and r.headers["content-type"].startswith("text/html")


def _contas(tmp):
    contas = ContasSQLite(tmp / "contas.db")
    contas.garantir_padrao()
    agora = "2026-10-21T00:00:00+00:00"
    grupos = {g["nome"]: g["grupo_id"] for g in contas.grupos()}
    so_chamadas = contas.criar_grupo("Só chamadas", agora)
    contas.definir_permissoes(so_chamadas, {"ver_painel", "ver_chamadas"})
    for login, grupo in (("revisor", grupos["Revisor"]), ("leitor", grupos["Leitor"]), ("chamadas", so_chamadas)):
        contas.criar_usuario(login, login, grupo, A.hash_senha(SENHA), agora)
    return contas


def _entrar(cli, login):
    csrf = re.search(r'name="csrf" value="([^"]+)"', cli.get("/entrar").text).group(1)
    r = cli.post("/entrar", data={"login": login, "senha": SENHA, "proximo": "/", "csrf": csrf}, follow_redirects=False)
    assert r.status_code == 303, r.text[:300]
    return cli


@pytest.fixture(scope="module")
def com_login(demo):
    return lambda login: _entrar(_cliente(demo, autenticacao=True, contas=_contas_modulo(demo)), login)


_CONTAS = {}


def _contas_modulo(demo):
    if "c" not in _CONTAS:
        _CONTAS["c"] = _contas(demo[2])
    return _CONTAS["c"]


def test_revisor_baixa_e_ve_o_botao(com_login):
    cli = com_login("revisor")
    assert cli.get("/execucoes/resultado.xlsx").status_code == 200
    assert 'href="/execucoes/resultado.xlsx"' in cli.get("/execucoes").text


def test_grupo_sem_os_nomes_dos_revisores_nao_baixa_nem_ve_o_botao(com_login):
    cli = com_login("chamadas")
    r = cli.get("/execucoes/resultado.xlsx")
    assert r.status_code == 403 and "Seu grupo não tem permissão para esta página ou ação." in r.text
    pagina = cli.get("/execucoes")
    assert pagina.status_code == 200 and "/execucoes/resultado.xlsx" not in pagina.text


def test_grupo_sem_ver_chamadas_nao_baixa(com_login):
    r = com_login("leitor").get("/execucoes/resultado.xlsx", follow_redirects=False)
    assert r.status_code == 403 and "spreadsheetml" not in r.headers.get("content-type", "")


def test_visitante_anonimo_nao_baixa(demo):
    cli = _cliente(demo, autenticacao=True, contas=_contas_modulo(demo))
    r = cli.get("/execucoes/resultado.xlsx", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/entrar?proximo=%2Fexecucoes%2Fresultado.xlsx"


def test_sem_execucoes_a_pagina_nao_mostra_o_botao(tmp_path):
    banco = tmp_path / "vazio.db"
    RepositorioSQLite(banco).fechar()
    cli = TestClient(criar_app(LeituraPainelSQLite(banco), lambda: RepositorioSQLite(banco), BrutoLocal(tmp_path),
                               _contexto(tmp_path, False)))
    assert "/execucoes/resultado.xlsx" not in cli.get("/execucoes").text
    r = cli.get("/execucoes/resultado.xlsx")
    assert r.status_code == 200                                             # o arquivo sai, com as abas vazias
    abas = _abas_xlsx(r.content)
    assert len(abas) == 62
    # so o cabecalho, menos as views que sempre tem linhas (contagens zeradas) e o resumo com as 27 UFs
    assert {n: len(l) - 1 for n, l in abas.items() if len(l) > 1} == {"PRISMA": 20, "Eficiencia": 1, "Auditoria_negativos": 1,
                                                                    "Resumo_UF": 27}
